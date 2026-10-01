"""
Загрузка каталога книгой Excel вместе с фотографиями на Django
(savdex/data/workbook.py) — все случаи теста Laravel
tests/Feature/Admin/ListingWorkbookImportTest.php, тем же порядком.

Книга собирается здесь же, а не лежит файлом: тест должен ломаться при
изменении разбора, а не при пересохранении фикстуры. Значения пишет
openpyxl, фотографии — openpyxl.drawing.image с настоящим PNG от Pillow
(загрузка проверяет файлы). Данные заводят фабрики tests/factories.py,
загрузка идёт в отдельном процессе Django — под ролью savdex_django,
если задан SAVDEX_PARITY_DJANGO_URL, как на сервере. Публичный диск —
во временном каталоге (Storage::fake у Laravel).

Нужен PostgreSQL (SAVDEX_PARITY_PG_URL); общая часть — в pg_admin.py.
Разбор рисунков, как их пишет Excel (относительные ссылки, twoCellAnchor) и
образец книги проверяются и без базы.
"""

from __future__ import annotations

import base64
import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any

import pytest

from . import factories
from .pg_admin import PYTHON, АДРЕС, ОКРУЖЕНИЕ, sql, свежая_база

HEADERS = ["Номер", "Название", "Компания", "Категория", "Цена", "Валюта", "Город", "Фото"]

#: Подсказка к строке без компании: как её загрузить
ВЫХОД = (
    ". Заполните столбец «Компания» или выберите «Компанию для строк без компании» в окне загрузки"
)

#: Шапка языкового листа: только переводимые поля
TEXT_HEADERS = ["Заголовок", "Описание", "Условия поставки", "Условия оплаты"]

КИРПИЧ = ["", "Кирпич керамический М150", "ООО «Стройбаза»", "", "", "", "", ""]
ПРЯЖА = ["", "Пряжа хлопковая 30/1", "ООО «Стройбаза»", "", "", "", "", ""]

#: Django в отдельном процессе: загрузка книг, итог — JSON-ом
ПРОБА = """
import base64, json, sys
import django
django.setup()
from savdex.data import workbook

a = json.loads(sys.stdin.read())
files = [(name, base64.b64decode(body)) for name, body in a["files"]]

if a["action"]:
    # Кнопка «Загрузить» в списке объявлений: от имени администратора
    out = workbook.import_workbooks(
        files, admin_id=a["admin"], replace=a["replace"], ip="203.0.113.7",
        default_company=a["company"], default_type=a["type"], publish=a["publish"],
    )
else:
    # ListingWorkbookImport::run, как его зовёт тест Laravel: без входа
    out = workbook.import_workbook(
        files[0][1], author_id=a["admin"], replace=a["replace"],
        default_company=a["company"], default_type=a["type"], publish=a["publish"],
    )

print(json.dumps(out, ensure_ascii=False))
"""


@pytest.fixture(scope="module")
def база() -> None:
    if not АДРЕС:
        pytest.skip("нет SAVDEX_PARITY_PG_URL — проверка требует PostgreSQL")

    свежая_база()


@pytest.fixture
def диск(база: None, tmp_path: Path) -> Path:
    """Чистые таблицы, администратор и пустой публичный диск."""
    sql(
        "truncate listings, listing_images, companies, categories, category_translations, "
        "cities, city_translations, countries, users, admin_actions restart identity cascade"
    )
    sql(
        "insert into users (name, email, password, is_admin, admin_role, status, created_at, "
        "updated_at) values ('Администратор', 'admin@savdex.uz', 'x', true, 'superadmin', "
        "'active', now(), now())"
    )

    return tmp_path


def _админ() -> int:
    return int(sql("select id from users where email = 'admin@savdex.uz'")[0][0])


def загрузить(
    диск: Path,
    *книги: bytes | tuple[str, bytes],
    replace: bool = False,
    кнопкой: bool = False,
    компания_по_умолчанию: int | None = None,
    тип: str = "supply",
    опубликовать: bool = False,
) -> dict[str, Any]:
    """
    Загрузить книги в Django; кнопкой — как действие админки (с журналом).
    компания_по_умолчанию и тип — поля окна загрузки.
    """
    files = [к if isinstance(к, tuple) else ("catalog.xlsx", к) for к in книги]
    вывод = subprocess.run(
        [sys.executable, "-c", ПРОБА],
        cwd=PYTHON,
        env={
            **ОКРУЖЕНИЕ,
            "DJANGO_SETTINGS_MODULE": "savdex.settings",
            "PYTHONPATH": str(PYTHON),
            "LARAVEL_ROOT": str(диск),
        },
        input=json.dumps(
            {
                "files": [(name, base64.b64encode(body).decode()) for name, body in files],
                "admin": _админ(),
                "replace": replace,
                "action": кнопкой,
                "company": компания_по_умолчанию,
                "type": тип,
                "publish": опубликовать,
            }
        ),
        capture_output=True,
        text=True,
    )
    assert вывод.returncode == 0, вывод.stderr[-3000:]

    return dict(json.loads(вывод.stdout))


def компания(name: str = "ООО «Стройбаза»") -> int:
    return factories.компания(name=name)


def товар(company: int | None, title: str, **поля: Any) -> int:
    """ListingFactory у компании; без компании — своя у каждого."""
    if company is not None:
        поля["company_id"] = company

    return factories.объявление(title=title, **поля)


def категория(name: str, parent: int | None = None) -> int:
    return factories.категория(name, parent_id=parent)


def перевод_категории(category: int, locale: str, name: str) -> None:
    sql(
        "insert into category_translations (category_id, locale, name, created_at, updated_at) "
        "values (%s, %s, %s, now(), now())",
        [category, locale, name],
    )


def город(name: str, *, locale: str = "ru") -> int:
    """Город с русским названием — как city() в тесте Laravel."""
    [(city,)] = sql(
        "insert into cities (country_id, slug, sort, is_active, created_at, updated_at) "
        "values (%s, 'tashkent', 0, true, now(), now()) returning id",
        [factories.узбекистан()],
    )
    перевод_города(city, locale, name)

    return int(city)


def перевод_города(city: int, locale: str, name: str) -> None:
    sql(
        "insert into city_translations (city_id, locale, name, created_at, updated_at) "
        "values (%s, %s, %s, now(), now())",
        [city, locale, name],
    )


def строки(query: str, params: list[Any] | None = None) -> list[dict[str, Any]]:
    import psycopg
    from psycopg.rows import dict_row

    with psycopg.connect(АДРЕС, autocommit=True, row_factory=dict_row) as соединение:
        return list(соединение.execute(query, params or []).fetchall())


def объявление(title: str) -> dict[str, Any]:
    [found] = строки("select * from listings where title = %s", [title])

    return found


def по_номеру(listing_id: int) -> dict[str, Any]:
    [found] = строки("select * from listings where id = %s", [listing_id])

    return found


def фото(listing_id: int) -> list[dict[str, Any]]:
    return строки(
        "select * from listing_images where listing_id = %s order by sort, id", [listing_id]
    )


def перевод(listing: dict[str, Any], field: str, locale: str) -> str | None:
    """Listing::hasTranslation / localized*: перевод поля, если он не пустой."""
    text = (listing[f"{field}_i18n"] or {}).get(locale)

    return text if text is not None and str(text).strip() != "" else None


# ── Сборка книги ─────────────────────────────────────────────────────


def png(seed: int) -> bytes:
    """Настоящий PNG: загрузка проверяет файлы."""
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (40, 30), ((seed * 40) % 255, 120, 200)).save(out, "PNG")

    return out.getvalue()


def книга(
    rows: dict[int, list[str]],
    pictures: dict[int, int],
    translations: dict[str, dict[int, list[str]]] | None = None,
    *,
    russian_last: bool = False,
    hidden: tuple[str, ...] = (),
) -> bytes:
    """
    Русский лист с товарами и, если нужно, языковые листы. Русский без
    языковых остаётся безымянным, как книга из одного листа; с ними —
    «Русский», первым или последним. Фотографии — на русский лист, в
    столбец «Фото» и правее, по одной на колонку.
    """
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.utils import get_column_letter

    translations = translations or {}
    book = Workbook()
    first = True

    def sheet(name: str | None, headers: list[str], body: dict[int, list[str]]) -> Any:
        nonlocal first
        page: Any = book.active if first else book.create_sheet()
        first = False

        if name is not None:
            page.title = name

            if name in hidden:
                page.sheet_state = "hidden"

        page.append(headers)

        for number in range(2, (max(body) if body else 1) + 1):
            page.append(body.get(number, [""]))

        return page

    def russian() -> Any:
        return sheet("Русский" if translations else None, HEADERS, rows)

    page = None if russian_last else russian()

    for name, body in translations.items():
        sheet(name, TEXT_HEADERS, body)

    if russian_last:
        page = russian()

    number = 0

    for row, count in pictures.items():
        for i in range(count):
            number += 1
            page.add_image(Image(io.BytesIO(png(number))), f"{get_column_letter(8 + i)}{row}")

    out = io.BytesIO()
    book.save(out)

    return out.getvalue()


def лист(name: str, rows: list[list[Any]]) -> bytes:
    """Книга из одного листа — со строками как есть, без готовой шапки."""
    from openpyxl import Workbook

    book = Workbook()
    page: Any = book.active
    page.title = name

    for row in rows:
        page.append(row or [""])

    out = io.BytesIO()
    book.save(out)

    return out.getvalue()


def рисунки_как_в_excel(content: bytes, pictures: dict[int, int], sheet: int = 1) -> bytes:
    """
    addPictures() теста Laravel: рисунки дописываются в готовую книгу так,
    как это делает Excel, — twoCellAnchor и ссылки относительно части.
    """
    anchors = relations = ""
    media: dict[str, bytes] = {}
    number = 0

    for row, count in pictures.items():
        for i in range(count):
            number += 1
            media[f"xl/media/image{number}.png"] = png(number)
            anchors += (
                '<xdr:twoCellAnchor editAs="oneCell">'
                f"<xdr:from><xdr:col>{7 + i}</xdr:col><xdr:colOff>0</xdr:colOff>"
                f"<xdr:row>{row - 1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>"
                f"<xdr:to><xdr:col>{8 + i}</xdr:col><xdr:colOff>0</xdr:colOff>"
                f"<xdr:row>{row}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>"
                f'<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="{number}" name="Picture {number}"/>'
                "<xdr:cNvPicPr/></xdr:nvPicPr>"
                '<xdr:blipFill><a:blip xmlns:r="http://schemas.openxmlformats.org/'
                f'officeDocument/2006/relationships" r:embed="rId{number}"/>'
                "<a:stretch><a:fillRect/></a:stretch></xdr:blipFill>"
                '<xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>'
                "</xdr:pic><xdr:clientData/></xdr:twoCellAnchor>"
            )
            relations += (
                f'<Relationship Id="rId{number}" Type="http://schemas.openxmlformats.org/'
                'officeDocument/2006/relationships/image" '
                f'Target="../media/image{number}.png"/>'
            )

    head = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    parts = {
        **media,
        "xl/drawings/drawing1.xml": head
        + '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/'
        'spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        + anchors
        + "</xdr:wsDr>",
        "xl/drawings/_rels/drawing1.xml.rels": head
        + '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + relations
        + "</Relationships>",
        f"xl/worksheets/_rels/sheet{sheet}.xml.rels": head
        + '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        'relationships/drawing" Target="../drawings/drawing1.xml"/></Relationships>',
    }

    return _дописать(content, parts)


def _дописать(content: bytes, parts: dict[str, str | bytes]) -> bytes:
    """ZipArchive::addFromString: части заменяются или добавляются."""
    out = io.BytesIO()

    with (
        zipfile.ZipFile(io.BytesIO(content)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            if item.filename not in parts:
                target.writestr(item, source.read(item.filename))

        for name, data in parts.items():
            target.writestr(name, data)

    return out.getvalue()


# ── Фотографии ───────────────────────────────────────────────────────


def test_фотографии_достаются_объявлениям_своих_строк(диск):
    c = компания()
    brick = товар(c, "Кирпич керамический М150")
    yarn = товар(c, "Пряжа хлопковая 30/1")

    # Между товарами пустая строка: номера строк не должны сползти
    result = загрузить(диск, книга({2: КИРПИЧ, 3: [], 4: ПРЯЖА}, {2: 1, 4: 1}))

    assert result["photos"] == 2
    assert len(фото(brick)) == 1 and len(фото(yarn)) == 1

    [image] = фото(brick)
    assert (диск / "storage/app/public" / image["path"]).is_file()
    assert (диск / "storage/app/public" / image["thumb_path"]).is_file()
    assert image["path"].startswith(f"listings/{brick}/")
    assert image["thumb_path"].startswith(f"listings/{brick}/thumb/")


def test_несколько_фотографий_строки_идут_по_порядку_колонок(диск):
    listing = товар(компания(), "Кирпич керамический М150")

    загрузить(диск, книга({2: КИРПИЧ}, {2: 3}))

    assert [image["sort"] for image in фото(listing)] == [0, 1, 2]


def test_новый_товар_создаётся_с_компанией_категорией_и_ценой(диск):
    c = компания()
    section = категория("Стройматериалы")

    result = загрузить(
        диск,
        книга(
            {
                2: [
                    "",
                    "Профнастил С8",
                    "ООО «Стройбаза»",
                    "Стройматериалы",
                    "1 200 000",
                    "сум",
                    "",
                    "",
                ]
            },
            {2: 1},
        ),
    )
    listing = объявление("Профнастил С8")

    assert result["created"] == 1
    assert listing["category_id"] == section and listing["company_id"] == c
    assert float(listing["price"]) == 1_200_000.0 and listing["currency"] == "UZS"
    assert listing["slug"] == f"profnastil-s8-{listing['id']}"
    assert len(фото(listing["id"])) == 1

    # Новое ждёт проверки: публикует администратор, посмотрев, что получилось
    assert listing["status"] == "moderation" and listing["published_at"] is None
    assert listing["source"] == "import" and listing["type"] == "supply"
    assert listing["user_id"] == _админ() and listing["price_negotiable"] is False
    assert "profnastil" in listing["search_text"]


def test_повторная_загрузка_не_плодит_фотографии(диск):
    listing = товар(компания(), "Кирпич керамический М150")

    загрузить(диск, книга({2: КИРПИЧ}, {2: 1}))
    second = загрузить(диск, книга({2: КИРПИЧ}, {2: 1}))

    assert second["photos"] == 0
    assert len(фото(listing)) == 1


def test_галочка_замены_меняет_фотографии(диск):
    listing = товар(компания(), "Кирпич керамический М150")

    загрузить(диск, книга({2: КИРПИЧ}, {2: 1}))
    [old] = фото(listing)

    загрузить(диск, книга({2: КИРПИЧ}, {2: 2}), replace=True)

    assert len(фото(listing)) == 2
    assert old["path"] not in [image["path"] for image in фото(listing)]
    assert not (диск / "storage/app/public" / old["path"]).exists()
    assert not (диск / "storage/app/public" / old["thumb_path"]).exists()


# ── Строка ───────────────────────────────────────────────────────────


def test_незнакомая_категория_пропускается_а_товар_загружается(диск):
    """
    Непонятная ячейка пропускается, а товар загружается: модератор
    дочинит категорию; пропуск записан в отчёт.
    """
    компания()

    result = загрузить(
        диск,
        книга({2: ["", "Профнастил С8", "ООО «Стройбаза»", "Строительство", "", "", "", ""]}, {}),
    )

    assert result["created"] == 1 and result["errors"] == []
    assert result["notes"] == [
        "Строка 2: категория «Строительство» не найдена в каталоге — ячейка пропущена"
    ]

    [listing] = строки("select * from listings")
    assert listing["title"] == "Профнастил С8" and listing["category_id"] is None


def test_строка_с_номером_правит_существующее_объявление(диск):
    listing = товар(компания(), "Старое название", price=100)

    result = загрузить(
        диск,
        книга({2: [str(listing), "Кирпич керамический М150", "", "", "250000", "", "", ""]}, {}),
    )
    fresh = по_номеру(listing)

    assert result["updated"] == 1
    assert fresh["title"] == "Кирпич керамический М150" and float(fresh["price"]) == 250_000.0

    # Опубликованное повторная загрузка не снимает с витрины
    assert fresh["status"] == "active" and fresh["source"] == "import"
    assert "kirpich" in fresh["search_text"]


def test_номер_числом_и_номер_которого_нет(диск):
    """
    «Номер» из Excel приходит дробным (12.0) — это всё равно номер 12.
    Номера нет в базе — строка не пропадает: загружается по заголовку
    и компании, как без номера.
    """
    listing = товар(компания(), "Кирпич керамический М150")
    book = книга(
        {
            2: [float(listing), "Кирпич новый", "", "", "", "", "", ""],
            3: [99999, "Нет такого", "ООО «Стройбаза»"],
        },
        {},
    )

    result = загрузить(диск, book)

    assert (result["updated"], result["created"], result["errors"]) == (1, 1, [])
    assert result["notes"] == [
        "Строка 3: объявление № 99999 не найдено — строка загружена по заголовку и компании"
    ]
    assert по_номеру(listing)["title"] == "Кирпич новый"
    assert объявление("Нет такого")["id"] != 99999


# ── Кнопка в админке ─────────────────────────────────────────────────


def test_кнопка_в_админке_принимает_книгу_и_раскладывает_фотографии(диск):
    listing = товар(компания(), "Кирпич керамический М150")

    result = загрузить(диск, книга({2: КИРПИЧ}, {2: 1}), кнопкой=True)

    assert (result["updated"], result["photos"], result["errors"]) == (1, 1, [])
    assert len(фото(listing)) == 1

    # Журнал: правка объявления (AuditObserver) и след загрузки пачкой
    [(subject, changes, ip)] = sql(
        "select subject_id, changes::text, ip from admin_actions "
        "where section = 'listings' and action = 'updated'"
    )
    assert (subject, ip) == (listing, "203.0.113.7")
    assert json.loads(changes)["after"] == {"source": "import"}
    assert sql(
        "select user_name, note, subject_id from admin_actions where action = 'imported'"
    ) == [("Администратор", "Книг: 1, создано: 0, обновлено: 1, фотографий: 1", None)]


def test_за_раз_принимается_несколько_книг(диск):
    c = компания()
    brick = товар(c, "Кирпич керамический М150")
    yarn = товар(c, "Пряжа хлопковая 30/1")

    # Каталог разрезан по разделам: в каждой книге свой товар
    result = загрузить(
        диск,
        ("stroy.xlsx", книга({2: КИРПИЧ}, {2: 1})),
        ("tekstil.xlsx", книга({2: ПРЯЖА}, {2: 1})),
        кнопкой=True,
    )

    assert (result["updated"], result["photos"]) == (2, 2)
    assert len(фото(brick)) == 1 and len(фото(yarn)) == 1
    assert sql("select note from admin_actions where action = 'imported'") == [
        ("Книг: 2, создано: 0, обновлено: 2, фотографий: 2",)
    ]


def test_имя_книги_в_отчёте_когда_книг_несколько(диск):
    """Из какой книги строка отчёта — понятно только когда их несколько."""
    компания()
    ошибка = книга({2: ["", "Без компании", "", "", "", "", "", ""]}, {})

    одна = загрузить(диск, ("one.xlsx", ошибка), кнопкой=True)
    две = загрузить(диск, ("one.xlsx", ошибка), ("two.xlsx", лист("Лист1", [])), кнопкой=True)

    assert одна["errors"] == [
        "Строка 2: не указана компания, а без неё новое объявление не создать" + ВЫХОД
    ]
    assert две["errors"] == [
        "one.xlsx, Строка 2: не указана компания, а без неё новое объявление не создать" + ВЫХОД,
        "two.xlsx, Лист «Лист1»: не найдено ни одной заполненной ячейки — лист пуст. "
        "Скачайте образец книги — в нём столбцы названы так, как их ждёт загрузка.",
    ]


# ── Образец и языки столбцов ─────────────────────────────────────────


def test_образец_книги_загружается_без_правок(диск):
    """Образец отдаём заказчику: если его же загрузка не читает — читать нечего."""
    from savdex.data.workbook import template_bytes

    компания()
    section = категория("Стройматериалы")
    bricks = категория("Кирпич и блоки", section)
    city = город("Ташкент")

    result = загрузить(диск, template_bytes())
    listing = объявление("Кирпич керамический М150")

    assert (result["errors"], result["notes"], result["created"]) == ([], [], 1)
    assert listing["category_id"] == bricks and listing["city_id"] == city
    assert listing["currency"] == "UZS" and float(listing["price"]) == 1_200.0
    assert listing["status"] == "moderation" and listing["unit"] == "шт"
    assert listing["min_order"] == 5000
    assert listing["delivery_terms"] == "Самовывоз со склада, доставка по Ташкентской области."
    assert listing["payment_terms"] == "Предоплата 50 %, остаток по факту отгрузки."

    # Четыре языковых листа образца дают четыре перевода каждого поля
    for locale in ("en", "uz", "zh", "tr"):
        for field in ("title", "description", "delivery_terms", "payment_terms"):
            assert перевод(listing, field, locale) is not None, (field, locale)

    assert перевод(listing, "title", "en") == "Ceramic brick M150"


def test_столбцы_на_любом_языке_узнаются(диск):
    # Каждый столбец назван на своём языке, значения — тоже
    компания("Uyut Gulistan Mebel")
    furniture = категория("Мебель")
    перевод_категории(furniture, "zh", "家具")
    city = город("Ташкент")
    перевод_города(city, "uz", "Toshkent")

    result = загрузить(
        диск,
        лист(
            "Sheet1",
            [
                ["Nomi", "Şirket", "类别", "Price", "Валюта", "Shahar", "Yayınla"],
                [
                    "Ofis stoli",
                    "Uyut Gulistan Mebel",
                    "家具",
                    "1 200 000",
                    "so'm",
                    "Toshkent",
                    "evet",
                ],
            ],
        ),
    )
    listing = объявление("Ofis stoli")

    assert result["errors"] == [] and result["created"] == 1
    assert listing["category_id"] == furniture and listing["city_id"] == city
    assert float(listing["price"]) == 1_200_000.0 and listing["currency"] == "UZS"
    assert listing["status"] == "moderation"


# ── Языковые листы ───────────────────────────────────────────────────


def test_языковые_листы_дают_переводы_строка_к_строке(диск):
    компания()

    # Два товара через пустую строку: смещение от шапки должно сходиться
    # на всех листах, включая пустые строки
    result = загрузить(
        диск,
        книга(
            {2: КИРПИЧ, 3: [], 4: ПРЯЖА},
            {2: 1},
            {
                "English": {
                    2: [
                        "Ceramic brick M150",
                        "Solid brick.",
                        "Pick-up from the warehouse",
                        "50 % prepayment",
                    ],
                    3: [],
                    4: ["Cotton yarn 30/1", "Combed yarn.", "", ""],
                },
                "中文": {
                    2: ["M150 陶瓷砖", "实心砖。", "仓库自提", "预付 50 %"],
                    3: [],
                    4: ["棉纱 30/1", "", "", ""],
                },
            },
        ),
    )
    brick = объявление("Кирпич керамический М150")
    yarn = объявление("Пряжа хлопковая 30/1")

    assert result["errors"] == [] and result["created"] == 2
    assert перевод(brick, "title", "en") == "Ceramic brick M150"
    assert перевод(brick, "description", "en") == "Solid brick."
    assert перевод(brick, "delivery_terms", "en") == "Pick-up from the warehouse"
    assert перевод(brick, "payment_terms", "en") == "50 % prepayment"
    assert перевод(brick, "title", "zh") == "M150 陶瓷砖"
    assert len(фото(brick["id"])) == 1

    assert перевод(yarn, "title", "en") == "Cotton yarn 30/1"
    assert перевод(yarn, "title", "zh") == "棉纱 30/1"

    # Чего на листе не было, того и в переводах нет
    assert перевод(yarn, "description", "zh") is None
    assert перевод(yarn, "delivery_terms", "en") is None
    assert перевод(brick, "title", "uz") is None

    # Переводы заголовка — в поисковом тексте
    assert "ceramic brick" in brick["search_text"]


def test_русский_лист_узнаётся_по_имени_а_не_по_порядку(диск):
    listing = товар(компания(), "Кирпич керамический М150")

    # Английский лист первый: цены и фото всё равно — с русского
    result = загрузить(
        диск,
        книга(
            {2: ["", "Кирпич керамический М150", "ООО «Стройбаза»", "", "9 900", "сум", "", ""]},
            {2: 1},
            {"English": {2: ["Ceramic brick M150", "", "", ""]}},
            russian_last=True,
        ),
    )
    fresh = по_номеру(listing)

    assert result["errors"] == [] and result["updated"] == 1
    assert len(фото(listing)) == 1
    assert float(fresh["price"]) == 9_900.0
    assert перевод(fresh, "title", "en") == "Ceramic brick M150"


def test_лист_с_другим_числом_строк_пропускается_а_товары_грузятся(диск):
    компания()

    # На английском листе строка удалена: переводы съехали бы на соседний
    # товар — лист пропускается, товары загружаются без английского
    result = загрузить(
        диск,
        книга({2: КИРПИЧ, 3: ПРЯЖА}, {}, {"English": {2: ["Cotton yarn 30/1", "", "", ""]}}),
    )

    assert (result["created"], result["errors"]) == (2, [])
    [note] = result["notes"]
    assert note.startswith("Лист «English»: строк с данными 1, на русском листе 2.")
    assert "лист пропущен" in note
    assert перевод(объявление("Кирпич керамический М150"), "title", "en") is None


def test_пустой_языковой_лист_не_мешает_загрузке(диск):
    компания()

    result = загрузить(диск, книга({2: КИРПИЧ}, {}, {"English": {}}))

    assert result["errors"] == [] and result["created"] == 1
    assert result["notes"] == ["Лист «English» пуст — переводов на этот язык нет."]

    [listing] = строки("select * from listings")
    assert перевод(listing, "title", "en") is None


def test_пустая_ячейка_перевода_не_стирает_старый(диск):
    listing = товар(
        компания(),
        "Кирпич керамический М150",
        title_i18n={"en": "Ceramic brick M150", "tr": "Seramik tuğla M150"},
    )

    # Переводчик дошёл только до описания: заголовок в книге пуст
    загрузить(диск, книга({2: КИРПИЧ}, {}, {"English": {2: ["", "Solid brick.", "", ""]}}))
    fresh = по_номеру(listing)

    assert перевод(fresh, "title", "en") == "Ceramic brick M150"
    assert перевод(fresh, "title", "tr") == "Seramik tuğla M150"
    assert перевод(fresh, "description", "en") == "Solid brick."


def test_второй_лист_на_том_же_языке_пропускается(диск):
    компания()

    result = загрузить(
        диск,
        книга(
            {2: КИРПИЧ},
            {},
            {"English": {2: ["Ceramic brick", "", "", ""]}, "en": {2: ["Brick", "", "", ""]}},
        ),
    )

    assert (result["created"], result["errors"]) == (1, [])
    assert result["notes"] == [
        "Два листа на одном языке: «English» и «en». Взят первый, «en» пропущен."
    ]
    assert перевод(объявление("Кирпич керамический М150"), "title", "en") == "Ceramic brick"


def test_книга_без_русского_листа_грузится_с_первого(диск):
    компания()

    result = загрузить(
        диск,
        лист("English", [HEADERS, ["", "Ceramic brick", "ООО «Стройбаза»", "", "", "", "", ""]]),
    )

    assert (result["created"], result["errors"]) == (1, [])
    assert result["notes"] == [
        "В книге нет листа «Русский» — главным взят первый лист «English»: цены, категории "
        "и фотографии берутся с него."
    ]


# ── Шапка ────────────────────────────────────────────────────────────


def test_шапка_с_пометками_узнаётся(диск):
    компания()

    # Звёздочки, пояснения в скобках и единица через запятую
    result = загрузить(
        диск,
        лист(
            "Лист XLSX",
            [
                ["Название*", "Компания:", "Цена (за штуку)", "Валюта, код"],
                ["Профнастил С8", "ООО «Стройбаза»", "1 200 000", "сум"],
            ],
        ),
    )
    listing = объявление("Профнастил С8")

    assert result["errors"] == []
    assert float(listing["price"]) == 1_200_000.0 and listing["currency"] == "UZS"


def test_непризнанная_шапка_пропускается_а_столбцы_угадываются(диск):
    """
    Названия столбцов свои («Наименование позиции», «Сколько стоит») — не
    отказ: строка названий узнаётся по тому, что под «Сколько стоит» числа,
    и пропускается, а столбцы определяются по содержимому.
    """
    компания()

    result = загрузить(
        диск,
        лист(
            "Лист XLSX",
            [
                ["Каталог продукции"],
                [],
                ["Наименование позиции", "Производитель", "Сколько стоит"],
                ["Профнастил С8", "ООО «Стройбаза»", "1 200 000"],
            ],
        ),
    )

    assert (result["created"], result["errors"]) == (1, [])
    assert result["notes"] == [
        "Лист «Лист XLSX»: названия столбцов в строке 3 не узнаны — столбцы определены по "
        "содержимому: A — заголовок, B — компания, C — цена. Проверьте загруженное; чтобы не "
        "гадать, назовите столбцы, как в образце."
    ]
    assert str(объявление("Профнастил С8")["price"]).startswith("1200000")


def test_без_шапки_и_без_заголовков_отчёт_показывает_что_стоит_в_строке(диск):
    """Угадать нечего (одни числа) — прежнее объяснение, что лежит на листе."""
    result = загрузить(диск, лист("Лист XLSX", [["Каталог продукции"], [], ["1", "2", "3"]]))

    assert result["created"] == 0
    [error] = result["errors"]
    assert error.startswith(
        "Лист «Лист XLSX»: не найдена строка с названиями столбцов. "
        "В строке 3 записано «1», «2», «3»"
    )
    assert "Загрузка ждёт: Номер, Название, Компания" in error and "Фото" not in error


def test_без_шапки_столбцы_угадываются_по_содержимому(диск):
    """
    Книга, у которой потерялась строка с названиями столбцов: столбцы
    узнаются по тому, что в них лежит, и отчёт говорит, что чем стало.
    """
    своя = компания("SAVDEX Заявки")
    раздел = категория("Оборудование")
    станки = категория("Станки", раздел)
    описание = "Покупатель ищет поставщика оборудования, требования подробно указаны в заявке."

    result = загрузить(
        диск,
        лист(
            "Лист XLSX",
            [
                ["Заявки на покупку, 18.09"],
                [
                    "",
                    "Куплю станок токарный",
                    описание,
                    "Оборудование → Станки",
                    "ООО «Нет в справочнике»",
                    "1200",
                    "USD",
                    "Самовывоз",
                    "Аккредитив",
                ],
                [
                    "",
                    "Куплю пресс гидравлический",
                    описание + " Второй.",
                    "Оборудование → Прессы",
                    "ТОО «Тоже нет»",
                    "",
                    "",
                    "Доставка, рассматриваются варианты",
                    "По договорённости",
                ],
            ],
        ),
        компания_по_умолчанию=своя,
        тип="demand",
    )

    assert (result["created"], result["errors"]) == (2, [])
    assert result["notes"][0] == (
        "Лист «Лист XLSX»: нет строки с названиями столбцов — столбцы определены по "
        "содержимому: B — заголовок, C — описание, D — категория, E — компания, F — цена, "
        "G — валюта, H — условия поставки, I — условия оплаты. Проверьте загруженное; чтобы "
        "не гадать, добавьте над таблицей строку с названиями, как в образце."
    )

    станок = объявление("Куплю станок токарный")
    assert (станок["category_id"], станок["company_id"], станок["type"]) == (станки, своя, "demand")
    assert (станок["currency"], станок["delivery_terms"], станок["payment_terms"]) == (
        "USD",
        "Самовывоз",
        "Аккредитив",
    )
    assert станок["description"] == описание
    assert объявление("Куплю пресс гидравлический")["category_id"] == раздел


def test_одно_похожее_слово_в_строке_данных_не_шапка(диск):
    """«Доставка, варианты» похоже на столбец «Доставка», но строка — данные."""
    компания()

    result = загрузить(
        диск,
        лист(
            "Лист1",
            [
                ["", "Кирпич керамический", "Доставка, рассматриваются варианты"],
                ["Номер", "Название", "Компания"],
                ["", "Пряжа хлопковая", "ООО «Стройбаза»"],
            ],
        ),
    )

    assert (result["created"], result["errors"]) == (1, [])
    assert объявление("Пряжа хлопковая")["title"] == "Пряжа хлопковая"


def test_шапка_в_одной_ячейке_просит_разбить_по_столбцам(диск):
    result = загрузить(
        диск,
        лист(
            "Лист XLSX",
            [["Номер;Название;Компания;Цена"], [";Профнастил С8;ООО «Стройбаза»;1200000"]],
        ),
    )

    [error] = result["errors"]
    assert "в одной ячейке через «;»" in error


def test_пустой_лист_называется_пустым(диск):
    result = загрузить(диск, лист("Лист XLSX", []))

    [error] = result["errors"]
    assert "лист пуст" in error


def test_условия_поставки_и_оплаты_читаются_с_русского_листа(диск):
    компания()

    result = загрузить(
        диск,
        лист(
            "Sheet1",
            [
                ["Название", "Компания", "Условия доставки", "Оплата"],
                [
                    "Кирпич керамический М150",
                    "ООО «Стройбаза»",
                    "Самовывоз со склада",
                    "Предоплата 50 %",
                ],
            ],
        ),
    )
    [listing] = строки("select * from listings")

    assert result["errors"] == []
    assert listing["delivery_terms"] == "Самовывоз со склада"
    assert listing["payment_terms"] == "Предоплата 50 %"


def test_незнакомая_валюта_пропускается(диск):
    """Валюта не из списка — цена остаётся в сумах, товар загружается."""
    компания()

    result = загрузить(
        диск,
        книга(
            {2: ["", "Кирпич керамический М150", "ООО «Стройбаза»", "", "100", "XYZ", "", ""]}, {}
        ),
    )

    assert result["created"] == 1 and result["errors"] == []
    assert "XYZ" in result["notes"][0]
    assert result["notes"] == [
        "Строка 2: валюта «XYZ» не поддерживается (можно: UZS, USD, EUR, CNY, TRY, RUB, KZT) "
        "— оставлена прежняя"
    ]

    [listing] = строки("select price::text, currency from listings")
    assert (listing["price"], listing["currency"]) == ("100.00", "UZS")


def test_скрытый_лист_не_читается(диск):
    компания()

    # «English» из образца спрятали вместо удаления: в нём остался пример,
    # и он не должен стать переводом чужого товара
    result = загрузить(
        диск,
        книга(
            {2: КИРПИЧ, 3: ПРЯЖА},
            {},
            {"English": {2: ["Ceramic brick M150", "Sample.", "", ""]}},
            hidden=("English",),
        ),
    )

    assert result["errors"] == [] and result["created"] == 2
    assert result["notes"] == ["Скрытый лист «English» пропущен."]
    assert перевод(объявление("Кирпич керамический М150"), "title", "en") is None


def test_лист_с_неузнанным_именем_пропускается_с_заметкой(диск):
    компания()

    result = загрузить(
        диск, книга({2: КИРПИЧ}, {}, {"Eng.": {2: ["Ceramic brick M150", "", "", ""]}})
    )

    assert result["errors"] == [] and result["created"] == 1
    [note] = result["notes"]
    assert "«Eng.»" in note
    [listing] = строки("select * from listings")
    assert перевод(listing, "title", "en") is None


def test_перевод_напротив_пустой_русской_строки_пропускает_лист(диск):
    компания()

    # Число строк сходится, но в русском листе строку стёрли, а внизу
    # дописали: переводы съехали на соседний товар
    result = загрузить(
        диск,
        книга(
            {2: КИРПИЧ, 3: [], 4: ПРЯЖА},
            {},
            {
                "English": {
                    2: ["Ceramic brick M150", "", "", ""],
                    3: ["Orphan translation", "", "", ""],
                    4: ["Cotton yarn 30/1", "", "", ""],
                }
            },
        ),
    )

    assert (result["created"], result["errors"]) == (2, [])
    assert result["notes"] == [
        "Лист «English», строка 3: заполнена, а на русском листе строка 3 пуста — строки "
        "разъехались. Лист пропущен, товары загружены без этого языка."
    ]
    assert перевод(объявление("Пряжа хлопковая 30/1"), "title", "en") is None


def test_одинаковый_заголовок_у_разных_компаний_требует_компанию(диск):
    first = товар(None, "Кирпич керамический М150", price=100)
    second = товар(None, "Кирпич керамический М150", price=200)

    # Ни номера, ни компании: править первое попавшееся — значит менять
    # цену чужого товара
    result = загрузить(
        диск, книга({2: ["", "Кирпич керамический М150", "", "", "777", "", "", ""]}, {})
    )

    assert result["updated"] == 0
    [error] = result["errors"]
    assert "нескольких объявлений" in error
    assert float(по_номеру(first)["price"]) == 100.0
    assert float(по_номеру(second)["price"]) == 200.0


def test_слишком_длинный_текст_обрезается_с_заметкой(диск):
    """Длинный текст обрезается по границе слова, а не отменяет строку."""
    компания()

    result = загрузить(диск, книга({2: КИРПИЧ}, {}, {"English": {2: ["Brick " * 20, "", "", ""]}}))

    assert result["created"] == 1 and result["errors"] == []
    assert "«en»" in result["notes"][0] and "90" in result["notes"][0]

    [listing] = строки("select * from listings")
    english = listing["title_i18n"]["en"]

    assert len(english) <= 90 and english.startswith("Brick Brick")
    assert english == ("Brick " * 15).strip()


# ── Сверх теста Laravel ──────────────────────────────────────────────


def test_пропуски_свои_у_каждой_строки(диск):
    """
    Заметка о пропущенной ячейке — только у своей строки. Раньше список
    не сбрасывался (ошибка Laravel, повторённая до его отключения), и
    заметка второй строки повторялась у всех следующих: на книге в
    двести строк — тысячи строк отчёта.
    """
    компания()

    result = загрузить(
        диск,
        книга(
            {
                2: ["", "Профнастил С8", "ООО «Стройбаза»", "Строительство", "", "", "", ""],
                3: ["", "Профнастил С10", "ООО «Стройбаза»", "", "", "", "", ""],
            },
            {},
        ),
    )

    assert result["created"] == 2
    assert result["notes"] == [
        "Строка 2: категория «Строительство» не найдена в каталоге — ячейка пропущена",
    ]


# ── Книга без компании и типа: заявки, собранные с других площадок ──


ЗАЯВКИ = ["Номер", "Заголовок", "Категория", "Компания", "Город", "Фото"]


def test_без_компании_строки_не_создаются_и_отчёт_подсказывает_выход(диск):
    result = загрузить(диск, лист("Русский", [ЗАЯВКИ, ["", "Куплю цемент М400", "", "", "", ""]]))

    assert result["created"] == 0
    assert result["errors"] == [
        "Строка 2: не указана компания, а без неё новое объявление не создать. Заполните "
        "столбец «Компания» или выберите «Компанию для строк без компании» в окне загрузки"
    ]


def test_компания_и_тип_из_окна_загрузки(диск):
    """
    Книга заявок на покупку: столбцов «Тип» нет, «Компания» пуста или
    называет продавца не из справочника. Строки достаются компании из
    окна загрузки и получают тип оттуда же.
    """
    своя = компания("SAVDEX Заявки")
    result = загрузить(
        диск,
        лист(
            "Русский",
            [
                ЗАЯВКИ,
                ["", "Куплю цемент М400", "", "", "", ""],
                ["", "Куплю арматуру А500", "", "ТОО «Нет в справочнике»", "", ""],
            ],
        ),
        компания_по_умолчанию=своя,
        тип="demand",
        кнопкой=True,
    )

    assert (result["created"], result["errors"]) == (2, [])
    assert result["notes"] == [
        "Строка 3: компания «ТОО «Нет в справочнике»» не найдена в справочнике — объявление "
        "отнесено к «SAVDEX Заявки»"
    ]

    for title in ("Куплю цемент М400", "Куплю арматуру А500"):
        строка = объявление(title)
        assert (строка["company_id"], строка["type"], строка["status"]) == (
            своя,
            "demand",
            "moderation",
        )


def test_без_компании_строки_уходят_служебной_компании(диск):
    """
    Компании в строке нет и в окне она не выбрана — заявка достаётся
    служебной компании Anjir Group (пока не найден настоящий владелец).
    """
    служебная = компания("OOO Anjir Group")
    result = загрузить(диск, лист("Русский", [ЗАЯВКИ, ["", "Куплю цемент М400", "", "", "", ""]]))

    assert (result["created"], result["errors"]) == (1, [])
    assert объявление("Куплю цемент М400")["company_id"] == служебная


def test_повторная_загрузка_с_компанией_передаёт_заявку_владельцу(диск):
    """Первая загрузка — без компании, вторая — с настоящей: та же заявка, без дубля."""
    служебная = компания("OOO Anjir Group")
    владелец = компания("ООО «Стройбаза»")
    строка = ["", "Куплю цемент М400", "", "", "", ""]

    загрузить(диск, лист("Русский", [ЗАЯВКИ, строка]))
    номер = объявление("Куплю цемент М400")["id"]
    assert по_номеру(номер)["company_id"] == служебная

    result = загрузить(
        диск, лист("Русский", [ЗАЯВКИ, ["", "Куплю цемент М400", "", "ООО «Стройбаза»", "", ""]])
    )

    assert (result["created"], result["updated"], result["errors"]) == (0, 1, [])
    assert result["notes"] == [
        "Строка 2: объявление передано от «OOO Anjir Group» компании «ООО «Стройбаза»»"
    ]
    assert по_номеру(номер)["company_id"] == владелец
    assert len(строки("select id from listings where deleted_at is null")) == 1


def test_галочка_сразу_опубликовать(диск):
    """
    Новые строки — сразу на витрине (с датой публикации и сроком);
    ждавшие проверки с прошлой загрузки — тоже; отклонённое — нет.
    """
    компания()
    книга_ = лист(
        "Русский",
        [
            ЗАЯВКИ,
            ["", "Куплю цемент М400", "", "ООО «Стройбаза»", "", ""],
            ["", "Куплю арматуру А500", "", "ООО «Стройбаза»", "", ""],
        ],
    )

    загрузить(диск, книга_)
    assert {r["status"] for r in строки("select status from listings")} == {"moderation"}
    sql("update listings set status = 'rejected' where title = 'Куплю арматуру А500'")

    result = загрузить(диск, книга_, опубликовать=True)

    assert (result["updated"], result["errors"]) == (2, [])
    цемент = объявление("Куплю цемент М400")
    assert цемент["status"] == "active"
    assert цемент["published_at"] is not None and цемент["expires_at"] is not None
    assert объявление("Куплю арматуру А500")["status"] == "rejected"

    новая = загрузить(
        диск,
        лист("Русский", [ЗАЯВКИ, ["", "Куплю кирпич М150", "", "ООО «Стройбаза»", "", ""]]),
        опубликовать=True,
    )
    assert новая["created"] == 1 and объявление("Куплю кирпич М150")["status"] == "active"


def test_компания_из_строки_важнее_компании_окна(диск):
    чужая = компания("ООО «Стройбаза»")
    своя = компания("SAVDEX Заявки")

    загрузить(
        диск,
        лист("Русский", [ЗАЯВКИ, ["", "Куплю кирпич", "", "ООО «Стройбаза»", "", ""]]),
        компания_по_умолчанию=своя,
    )

    строка = объявление("Куплю кирпич")
    assert (строка["company_id"], строка["type"]) == (чужая, "supply")


def test_компания_окна_не_переносит_существующее_объявление(диск):
    """Строка с «Номером» без компании правит объявление, продавец прежний."""
    прежняя = компания("ООО «Стройбаза»")
    своя = компания("SAVDEX Заявки")
    номер = товар(прежняя, "Кирпич керамический")

    result = загрузить(
        диск,
        лист("Русский", [ЗАЯВКИ, [номер, "Кирпич керамический М150", "", "", "", ""]]),
        компания_по_умолчанию=своя,
    )

    assert result["updated"] == 1
    assert по_номеру(номер)["company_id"] == прежняя


def test_прочее_это_другое_а_неизвестный_подраздел_это_раздел(диск):
    компания()
    раздел = категория("Оборудование")
    другое = категория("Другое", раздел)
    категория("Станки", раздел)
    # «Другое» есть и в соседнем разделе — выбирается своё
    категория("Другое", категория("Упаковка"))

    result = загрузить(
        диск,
        лист(
            "Русский",
            [
                ЗАЯВКИ,
                ["", "Куплю сепаратор", "Оборудование → Прочее", "ООО «Стройбаза»", "", ""],
                ["", "Куплю пресс", "Оборудование → Прессы", "ООО «Стройбаза»", "", ""],
            ],
        ),
    )

    assert result["created"] == 2
    assert объявление("Куплю сепаратор")["category_id"] == другое
    assert объявление("Куплю пресс")["category_id"] == раздел
    assert result["notes"] == [
        "Строка 3: подраздел «Прессы» не найден в каталоге — объявление в разделе «Оборудование»"
    ]


def test_повторная_загрузка_длинного_заголовка_не_заводит_дубль(диск):
    """Длинный заголовок хранится обрезанным — и ищется таким же."""
    компания()
    длинный = "Куплю " + "очень длинное название товара " * 5
    книга_ = лист("Русский", [ЗАЯВКИ, ["", длинный, "", "ООО «Стройбаза»", "", ""]])

    первая = загрузить(диск, книга_)
    вторая = загрузить(диск, книга_)

    assert (первая["created"], вторая["created"], вторая["updated"]) == (1, 0, 1)
    assert len(строки("select id from listings where deleted_at is null")) == 1


def test_рисунки_как_их_вставляет_excel(диск):
    """Относительные ссылки и twoCellAnchor, как в addPictures() теста Laravel."""
    listing = товар(компания(), "Кирпич керамический М150")

    result = загрузить(диск, рисунки_как_в_excel(книга({2: КИРПИЧ}, {}), {2: 2}))

    assert result["photos"] == 2
    assert [image["sort"] for image in фото(listing)] == [0, 1]


# ── Без базы ─────────────────────────────────────────────────────────


def test_фотографии_по_строкам_и_колонкам(tmp_path):
    """
    WorkbookImages: строка — по верхнему левому углу, в строке — по
    колонкам; absoluteAnchor, внешняя картинка и пустой файл — мимо.
    """
    from savdex.data.workbook import extract_photos

    book = рисунки_как_в_excel(книга({2: КИРПИЧ, 3: ПРЯЖА}, {}), {3: 2, 2: 1})

    with zipfile.ZipFile(io.BytesIO(book)) as source:
        drawing = source.read("xl/drawings/drawing1.xml").decode()
        rels = source.read("xl/drawings/_rels/drawing1.xml.rels").decode()

    # Колонки второй строки — в обратном порядке: порядок берётся из колонок
    drawing = drawing.replace("<xdr:col>7</xdr:col>", "<xdr:col>9</xdr:col>", 1)
    extra = (
        '<xdr:absoluteAnchor><xdr:pos x="0" y="0"/><xdr:pic><xdr:blipFill>'
        '<a:blip xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'r:embed="rId1"/></xdr:blipFill></xdr:pic></xdr:absoluteAnchor>'
        "<xdr:oneCellAnchor><xdr:from><xdr:col>0</xdr:col><xdr:row>3</xdr:row></xdr:from>"
        '<xdr:pic><xdr:blipFill><a:blip xmlns:r="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships" r:embed="rId9"/></xdr:blipFill></xdr:pic>'
        "</xdr:oneCellAnchor>"
        "<xdr:oneCellAnchor><xdr:from><xdr:col>0</xdr:col><xdr:row>4</xdr:row></xdr:from>"
        '<xdr:pic><xdr:blipFill><a:blip xmlns:r="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships" r:embed="rId8"/></xdr:blipFill></xdr:pic>'
        "</xdr:oneCellAnchor>"
    )
    drawing = drawing.replace("</xdr:wsDr>", extra + "</xdr:wsDr>")
    rels = rels.replace(
        "</Relationships>",
        '<Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        'relationships/image" Target="http://example.com/a.png" TargetMode="External"/>'
        '<Relationship Id="rId8" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        'relationships/image" Target="../media/empty.png"/></Relationships>',
    )
    book = _дописать(
        book,
        {
            "xl/drawings/drawing1.xml": drawing,
            "xl/drawings/_rels/drawing1.xml.rels": rels,
            "xl/media/empty.png": b"",
        },
    )

    photos = extract_photos(book, tmp_path, "Sheet")

    assert sorted(photos) == [2, 3]
    assert [p.name for p in photos[2]] == ["1-image3.png"]
    # Рисунок из колонки H перенесён в J: теперь он второй
    assert [p.name for p in photos[3]] == ["2-image2.png", "3-image1.png"]
    assert photos[3][1].read_bytes() == png(1)

    # Лист без рисунков и чужое имя — пусто; не книга — пусто
    assert extract_photos(book, tmp_path / "x", "Нет такого") == {}
    assert extract_photos(b"not a zip", tmp_path / "y") == {}


def test_рисунки_openpyxl_с_абсолютными_ссылками(tmp_path):
    from savdex.data.workbook import extract_photos

    photos = extract_photos(книга({2: КИРПИЧ, 3: [], 4: ПРЯЖА}, {4: 2, 2: 1}), tmp_path)

    assert {row: [p.name for p in paths] for row, paths in photos.items()} == {
        2: ["1-image3.png"],
        4: ["2-image1.png", "3-image2.png"],
    }


def test_образец_книги():
    """Пять листов, по одному на язык, с именами как в переключателе языка."""
    from openpyxl import load_workbook

    from savdex.data.workbook import HEADERS as ОБРАЗЕЦ
    from savdex.data.workbook import sheet_locale, template_bytes

    book = load_workbook(io.BytesIO(template_bytes()))
    names = book.sheetnames

    assert names == ["Русский", "English", "O‘zbekcha", "中文", "Türkçe"]
    assert [sheet_locale(name) for name in names] == ["ru", "en", "uz", "zh", "tr"]

    ru = [list(row) for row in book["Русский"].iter_rows(values_only=True)]
    assert ru[0] == ОБРАЗЕЦ and ru[1][0] is None and ru[1][5] == "1200"
    english = [list(row) for row in book["English"].iter_rows(values_only=True)]
    assert english[0] == ["Title", "Description", "Delivery terms", "Payment terms"]


def test_не_книга():
    from savdex.data.workbook import UnreadableWorkbookError, import_workbook

    with pytest.raises(UnreadableWorkbookError):
        import_workbook(b"not a workbook", author_id=1)


def test_значения_ячеек_как_у_openspout():
    from datetime import datetime, time

    from savdex.data.workbook import _strimwidth, _text

    epoch = datetime(1899, 12, 30)

    assert _text(12.0, epoch) == "12" and _text(12.5, epoch) == "12.5"
    assert _text(0.00001, epoch) == "1.0E-5"
    assert _text(True, epoch) == "да" and _text(False, epoch) == "нет"
    assert _text(datetime(2026, 10, 30, 12), epoch) == "30.10.2026"
    assert _text(time(12), epoch) == "30.12.1899"
    assert _text("  Кирпич \t", epoch) == "Кирпич" and _text(None, epoch) == ""
    # Неразрывный пробел trim() у PHP не трогает
    assert _text(" x ", epoch) == " x "

    assert _strimwidth("а" * 40, 40) == "а" * 40
    assert _strimwidth("а" * 41, 40) == "а" * 39 + "…"
    assert _strimwidth("中" * 25, 40) == "中" * 19 + "…"
