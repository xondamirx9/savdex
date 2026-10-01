"""
Выгрузка в Excel (manage.py export_xlsx) на настоящей базе с крайними
случаями.

Собственная сверка выгрузки перечитывает базу и сравнивает с книгой
лист за листом, ячейку за ячейкой. Слепой она была бы без крайних
случаев: с неверным чтением полей JSON сверка честно отвечала «все листы
сошлись». Поэтому база наполняется так, чтобы они были (наполнение() —
бывший tests/fixtures/export_parity_fill.php), а отдельная проверка
следит, что они действительно попали в книгу.

Порядок: схема — снимок миграций (свежая_база), справочники —
manage.py seed --fresh, затем наполнение и выгрузка.

Нужен PostgreSQL. Адрес базы — SAVDEX_PARITY_PG_URL; без неё проверка
пропускается.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path
from urllib.parse import urlparse

import pytest

from .factories import (
    компании,
    объявления,
    отзыв,
    открытие_контакта,
    пользователь,
    тендер,
)
from .pg_admin import ОКРУЖЕНИЕ, sql, свежая_база

PYTHON = Path(__file__).resolve().parents[1]
АДРЕС = os.environ.get("SAVDEX_PARITY_PG_URL", "")

pytestmark = pytest.mark.skipif(
    not АДРЕС,
    reason="нет SAVDEX_PARITY_PG_URL — выгрузка требует PostgreSQL",
)


def _проверочная_ли(url: str) -> bool:
    """
    База пересоздаётся с нуля — только та, что названа проверочной.

    Схема строится заново и стирает всё. Переменную окружения легко
    перепутать, и цена ошибки здесь — боевая база. Поэтому имя базы
    обязано содержать «test».
    """
    return "test" in urlparse(url).path.lstrip("/")


def _как_json_encode(value: object) -> str:
    """json_encode() по умолчанию: \\u-экранирование и «\\/» вместо «/»."""
    return json.dumps(value, separators=(",", ":")).replace("/", "\\/")


def наполнение() -> None:
    """
    Крайние случаи, на которых выгрузка могла бы ошибиться:

    - JSON с экранированной кириллицей (так пишет json_encode у PHP) и
      JSON с «неровными» пробелами — в книгу он идёт сырым текстом;
    - текст длиннее предела ячейки Excel — обрезка с пометкой;
    - ИНН и телефоны с ведущими нулями — остаются строками;
    - удалённые записи — выгружаются намеренно;
    - пустые значения, логические поля, дробные координаты.
    """
    [(plan,)] = sql("select id from plans order by id limit 1")
    cats = [c for (c,) in sql("select id from categories order by id")]
    companies = компании(12)

    for i, c in enumerate(companies):
        for _ in range(1 + i % 3):
            пользователь(company_id=c)

        объявления(i % 4, company_id=c)
        sql(
            "insert into wallets (company_id, credits, created_at, updated_at) "
            "values (%s, %s, now(), now())",
            [c, i * 3],
        )
        sql(
            "insert into subscriptions (company_id, plan_id, started_at, created_at, updated_at) "
            "values (%s, %s, now() - make_interval(days => %s), now(), now())",
            [c, plan, i],
        )
        sql(
            "insert into company_contacts (company_id, type, value, created_at, updated_at) "
            "values (%s, 'phone', %s, now(), now())",
            [c, f"00998{i:07d}"],
        )
        sql(
            "insert into company_attributes (company_id, key, value, created_at, updated_at) "
            "values (%s, 'сертификат', %s, now(), now())",
            [c, f"ISO 900{i}"],
        )

        if cats:
            sql(
                "insert into company_category (company_id, category_id) values (%s, %s)",
                [c, cats[i % len(cats)]],
            )

        sql(
            "insert into company_documents (company_id, type, title, file_path, created_at, "
            "updated_at) values (%s, 'license', %s, %s, now(), now())",
            [c, f"Лицензия №{i}", f"documents/{c}/l.pdf"],
        )

    for _ in range(4):
        тендер()

    listings = [pk for (pk,) in sql("select id from listings order by id")]

    for n, listing in enumerate(listings):
        sql(
            "insert into listing_images (listing_id, path, created_at, updated_at) "
            "values (%s, %s, now(), now())",
            [listing, f"listings/{listing}/a.webp"],
        )
        sql(
            "insert into listing_attributes (listing_id, key, value, created_at, updated_at) "
            "values (%s, 'Марка', %s, now(), now())",
            [listing, f"М{400 + n}"],
        )

        for d in range(3):
            sql(
                "insert into listing_stats (listing_id, date, views, created_at, updated_at) "
                "values (%s, current_date - %s, %s, now(), now())",
                [listing, d, n + d],
            )

    [(u,)] = sql("select id from users order by id limit 1")

    for listing in listings[:3]:
        sql(
            "insert into favorites (user_id, listing_id, created_at, updated_at) "
            "values (%s, %s, now(), now())",
            [u, listing],
        )

    # Раскрытия и отзывы — между разными компаниями
    for i in range(1, 5):
        открытие_контакта(company_id=companies[i], target_company_id=companies[0])
        отзыв(company_id=companies[0], author_company_id=companies[i], listing_id=None)

    # ── Крайние случаи ──
    sql(
        "update companies set description = %s, lat = 41.3110810, lng = 69.2405620, "
        "tin = '000123456', it_specializations = %s where id = %s",
        [
            "Длинное описание. " * 2500,  # > 32000 знаков — обрезка
            _как_json_encode(["веб", "мобильные/приложения"]),  # экранированный \\u и \\/
            companies[1],
        ],
    )
    sql("update companies set deleted_at = now() - interval '1 day' where id = %s", [companies[2]])
    sql(
        # json, записанный не через json_encode
        'update companies set it_specializations = \'["сырой текст",  "с пробелами"]\', '
        "website = null, phone = null where id = %s",
        [companies[3]],
    )
    sql(
        "update users set is_admin = true, admin_role = 'superadmin', admin_permissions = %s "
        "where id = (select id from users where company_id = %s order by id limit 1)",
        [_как_json_encode({"reports": ["view", "export"]}), companies[4]],
    )
    sql(
        "update listings set tags = %s, title_i18n = %s, deleted_at = now() where id = %s",
        [
            json.dumps(["цемент", "М400"], ensure_ascii=False, separators=(",", ":")),
            _как_json_encode({"en": "Cement", "uz": "Sement"}),
            listings[0],
        ],
    )

    # JSON в том виде, в каком его писал Laravel (\\u-экранирование), и в
    # произвольном: оба должны попасть в выгрузку как есть
    alive = [pk for (pk,) in sql("select id from listings where deleted_at is null order by id")]
    sql(
        "update listings set tags = %s where id = %s",
        [_как_json_encode(["цемент", "М400"]), alive[0]],
    )
    sql(
        "update listings set title_i18n = %s where id = %s",
        ['{"en":  "Cement \\/ bags",   "uz": "Sement"}', alive[1]],
    )


def выгрузить(каталог: Path, *args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "manage.py", "export_xlsx", f"--dir={каталог}", *args],
        cwd=PYTHON,
        env={**os.environ, "DATABASE_URL": АДРЕС},
        capture_output=True,
        text=True,
        check=True,
    )


@pytest.fixture(scope="module")
def выгрузки(tmp_path_factory):
    if not _проверочная_ли(АДРЕС):
        pytest.fail(
            "SAVDEX_PARITY_PG_URL ведёт в базу без «test» в имени. "
            "Проверка стирает базу целиком — отказываюсь."
        )

    свежая_база()
    # Справочники — владельцем базы, как при деплое
    subprocess.run(
        [sys.executable, "manage.py", "seed", "--fresh"],
        cwd=PYTHON,
        env={**ОКРУЖЕНИЕ, "DJANGO_DATABASE_URL": АДРЕС, "PYTHONPATH": str(PYTHON)},
        capture_output=True,
        check=True,
    )
    наполнение()

    py_dir = tmp_path_factory.mktemp("python")
    python = выгрузить(py_dir)

    return py_dir, python.stdout


def test_адрес_проверочной_базы_распознаётся():
    assert _проверочная_ли("postgres://u:p@h:5432/savdex_test")
    assert not _проверочная_ли("postgres://u:p@h:5432/savdex")


def test_выгрузка_сошлась_сама_с_собой(выгрузки):
    py_dir, python = выгрузки

    assert "Все листы сошлись" in python
    # Обе книги на месте
    assert next(py_dir.glob("savdex-companies-*.xlsx"))
    assert next(py_dir.glob("savdex-listings-*.xlsx"))


def test_крайние_случаи_действительно_в_выгрузке(выгрузки):
    """
    Проверка самой проверки.

    Без крайних случаев сверка проходила и с ошибкой в чтении JSON.
    Если наполнение однажды перестанет их создавать, сверка снова
    станет слепой — этот тест заметит это первым. JSON в книге — сырым
    текстом из базы, как его записали.
    """
    from openpyxl import load_workbook

    py_dir, _ = выгрузки
    listings = load_workbook(next(py_dir.glob("savdex-listings-*.xlsx")))["Объявления"]
    values = [str(cell.value) for row in listings.iter_rows() for cell in row if cell.value]

    assert any("\\u0446" in v for v in values), "нет JSON с экранированной кириллицей"
    assert any('"en":  "Cement' in v for v in values), "нет JSON с неровными пробелами"

    companies = load_workbook(next(py_dir.glob("savdex-companies-*.xlsx")))["Компании"]
    cells = [str(cell.value) for row in companies.iter_rows() for cell in row if cell.value]

    assert any(v.endswith("[…обрезано]") for v in cells), "нет обрезанного длинного текста"
    assert "000123456" in cells, "нет ИНН с ведущими нулями"


def test_режим_сверки_отдаёт_итог(выгрузки, tmp_path):
    """
    Режим сверки с готовыми книгами (--compare-with) и итогом для
    вызывающего (--json): строка с меткой SAVDEX-RESULT. Повторная
    выгрузка той же базы сходится с первой ячейка в ячейку.
    """
    py_dir, _ = выгрузки
    python = выгрузить(tmp_path, f"--compare-with={py_dir}", "--json")

    line = next(x for x in python.stdout.splitlines() if x.startswith("SAVDEX-RESULT "))
    result = json.loads(line.removeprefix("SAVDEX-RESULT "))

    assert result["self_check"] is True
    assert result["compared"] is True
    assert result["differences"] == 0, result["problems"]


def test_снимок_показывает_базу_на_момент_снимка(выгрузки, tmp_path):
    """
    --snapshot: Python-версия читает базу такой, какой её видел
    вызывающий в момент снимка, а не такой, какой она стала потом.

    Держим снимок открытым, как держит его админка, пока работает
    Python; тем временем «посетители» добавляют по тысяче просмотров
    каждому тендеру. В книге обязаны оказаться прежние числа.
    """
    import psycopg

    with psycopg.connect(АДРЕС) as снимок, psycopg.connect(АДРЕС, autocommit=True) as сайт:
        снимок.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY")
        идентификатор = снимок.execute("select pg_export_snapshot()").fetchone()[0]
        было = dict(снимок.execute("select id, views_count from tenders").fetchall())

        сайт.execute("update tenders set views_count = views_count + 1000")

        выгрузить(tmp_path, f"--snapshot={идентификатор}")

    assert было, "в базе нет тендеров — проверка слепа"
    assert _просмотры_тендеров(tmp_path) == было


def test_чужой_снимок_не_попадает_в_запрос(tmp_path):
    """
    SET TRANSACTION SNAPSHOT не принимает параметров запроса, значение
    вставляется в текст. Поэтому всё, что не похоже на снимок, —
    отказ до обращения к базе.
    """
    python = subprocess.run(
        [
            sys.executable,
            "manage.py",
            "export_xlsx",
            f"--dir={tmp_path}",
            "--snapshot=00000003-0000001B-1'; drop table tenders; --",
        ],
        cwd=PYTHON,
        env={**os.environ, "DATABASE_URL": АДРЕС},
        capture_output=True,
        text=True,
    )

    assert python.returncode == 1
    assert "Неверный снимок базы" in python.stderr
    assert list(tmp_path.iterdir()) == []


def _просмотры_тендеров(каталог: Path) -> dict[int, int]:
    from openpyxl import load_workbook

    for книга in каталог.glob("savdex-*.xlsx"):
        листы = load_workbook(книга, read_only=True)

        if "Тендеры" not in листы.sheetnames:
            continue

        строки = iter(листы["Тендеры"].iter_rows(values_only=True))

        for шапка in строки:
            if "Просмотров" in шапка:
                break

        ид, просмотры = шапка.index("ID"), шапка.index("Просмотров")

        return {int(r[ид]): int(r[просмотры]) for r in строки if r[ид] is not None}

    raise AssertionError("в книгах нет листа «Тендеры»")


def test_выгрузка_на_живой_базе_сходится(выгрузки, tmp_path):
    """
    Выгрузка с кнопки, пока посетители смотрят тендеры и объявления.

    Так и было на боевом сервере: две сверки из пяти «расходились» на
    одну ячейку — «Тендеры, Просмотров: в базе 40, в файле 39». Кто-то
    открыл тендер посреди выгрузки, и сверка сравнивала файл со
    сдвинувшейся базой. Файлы при этом были верны; врала сверка.

    Здесь «посетители» — поток, который без остановки увеличивает
    счётчики просмотров всё время, пока идёт выгрузка. Без единого
    снимка базы сверка расходится почти наверняка; со снимком — никогда.
    """
    import threading

    import psycopg

    стоп = threading.Event()
    обновлений = 0

    def посетители() -> None:
        nonlocal обновлений

        with psycopg.connect(АДРЕС, autocommit=True) as соединение:
            while not стоп.is_set():
                соединение.execute("update tenders set views_count = views_count + 1")
                соединение.execute("update listing_stats set views = views + 1")
                обновлений += 1

    поток = threading.Thread(target=посетители, daemon=True)
    поток.start()

    try:
        python = выгрузить(tmp_path, "--json")
    finally:
        стоп.set()
        поток.join(timeout=10)

    line = next(x for x in python.stdout.splitlines() if x.startswith("SAVDEX-RESULT "))
    result = json.loads(line.removeprefix("SAVDEX-RESULT "))

    assert обновлений > 10, "поток посетителей не успел ничего поменять — проверка слепа"
    assert result["self_check"] is True, result["problems"]
