"""
Выгрузка в Excel: приведение значений, описание листов, запись и сверка.

Части без PostgreSQL. Значения приводятся так же, как их видела
PHP-версия (книги, которыми пользуются, не меняются после переноса).
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from savdex.export.sheets import COMPANY_SHEETS, LISTING_SHEETS, Sheet
from savdex.export.workbooks import (
    CELL_LIMIT,
    Book,
    Collected,
    Exporter,
    compare,
    php_int,
    php_string,
    same,
    width,
    write,
)


class Заглушка:
    """Выгрузчик без базы: для проверки приведения значений."""

    vendor = "sqlite"


def выгрузчик(**dictionaries) -> Exporter:
    exporter = Exporter(Заглушка())  # type: ignore[arg-type]
    exporter.dictionaries = dictionaries

    return exporter


class TestКакВPHP:
    """Значения из psycopg приводятся к тому, что увидел бы PHP из PDO."""

    @pytest.mark.parametrize(
        ("value", "expected"),
        [
            (None, ""),
            (False, ""),
            (True, "1"),
            (0, "0"),
            (Decimal("4.50"), "4.50"),
            ("00998123", "00998123"),
            (datetime(2026, 9, 26, 11, 0), "2026-09-26 11:00:00"),
            (date(2026, 9, 26), "2026-09-26"),
        ],
    )
    def test_строка(self, value, expected):
        assert php_string(value) == expected

    def test_json_запасным_путём_как_json_encode(self):
        """
        Запасной путь, если json всё-таки пришёл словарём.

        PHP-версия кодирует с JSON_UNESCAPED_UNICODE: кириллица как есть,
        «/» экранирован, без пробелов.
        """
        assert php_string({"a": "веб/мобайл"}) == '{"a":"веб\\/мобайл"}'

    @pytest.mark.parametrize(
        ("value", "expected"),
        [(5, 5), ("12", 12), (Decimal("12.9"), 12), (Decimal("-12.9"), -12), (True, 1)],
    )
    def test_целое_отбрасывает_дробь_к_нулю(self, value, expected):
        assert php_int(value) == expected


class TestЯчейка:
    @pytest.mark.parametrize(
        ("value", "kind", "expected"),
        [
            (None, "int", ""),
            ("", "text", ""),
            (Decimal("7"), "int", 7),
            (Decimal("41.3110810"), "float", 41.311081),
            (True, "bool", "да"),
            (False, "bool", "нет"),
            (0, "bool", "нет"),
            ("f", "bool", "нет"),
            (datetime(2026, 9, 26, 11, 0, 5), "datetime", "2026-09-26 11:00:05"),
            ("2026-09-26T11:00:05.123Z", "datetime", "2026-09-26 11:00:05"),
            (datetime(2026, 9, 26, 11, 0), "date", "2026-09-26"),
            ("00998123", "text", "00998123"),
        ],
    )
    def test_приведение(self, value, kind, expected):
        assert выгрузчик().cell(value, kind) == expected

    def test_телефон_остаётся_строкой(self):
        """Excel превратил бы «00998…» в число и съел ведущие нули."""
        assert isinstance(выгрузчик().cell("00998901234567", "text"), str)

    def test_длинный_текст_обрезается_с_пометкой(self):
        exporter = выгрузчик()
        value = exporter.cell("ы" * (CELL_LIMIT + 10), "text")

        assert isinstance(value, str)
        assert value.endswith(" […обрезано]")
        assert exporter.truncated == 1


class TestНазвания:
    def test_название_вместо_идентификатора(self):
        exporter = выгрузчик(companies={3: "ООО «Цемент»"})

        assert exporter.lookup("@company", {"company_id": 3}) == "ООО «Цемент»"

    def test_пустой_идентификатор_даёт_пусто(self):
        assert выгрузчик(companies={}).lookup("@company", {"company_id": None}) == ""

    def test_счётчик_без_записей_это_ноль(self):
        """«Объявлений нет» и «не считали» читаются по-разному."""
        exporter = выгрузчик(listings_total={})

        assert exporter.lookup("@listings_total", {"id": 5}) == 0

    def test_неизвестное_поле_даёт_пусто(self):
        assert выгрузчик().lookup("@nonexistent", {"id": 1}) == ""


class TestСверка:
    def test_пустое_равно_пустому(self):
        assert same("", None)

    def test_число_из_excel_приходит_дробным(self):
        assert same(5, 5.0)

    def test_дробное_против_пустого_не_совпадает(self):
        assert not same(1.5, None)

    def test_расхождение_названо_с_адресом(self):
        problems = compare("Компании", ["ID", "Название"], [[1, "Цемент"]], [[1, "Песок"]])

        assert problems == ["Компании строка 2, «Название»: в базе [Цемент], в файле [Песок]"]

    def test_пропавшая_строка_видна_по_идентификатору(self):
        problems = compare("Компании", ["ID"], [[1], [2]], [[1], [3]])

        assert "в файле нет записей с ID 2" in problems[0]


class TestШирина:
    @pytest.mark.parametrize(
        ("header", "kind", "expected"),
        [
            ("ID компании", "int", 9),
            ("Рейтинг", "float", 14),
            ("Дата", "date", 12),
            ("Создана", "datetime", 19),
            ("Описание", "text", 60),
            ("Название", "text", 26),
        ],
    )
    def test_ширина_как_у_php_версии(self, header, kind, expected):
        assert width(header, kind) == expected


class TestОписаниеЛистов:
    """
    Описание листов: те же шестнадцать листов в двух книгах, что были у
    PHP-версии (ExportWorkbooks), и заголовки внутри листа различимы.
    """

    def test_листы_книг(self):
        assert [(s.name, s.table) for s in COMPANY_SHEETS] == [
            ("Компании", "companies"),
            ("Сотрудники", "users"),
            ("Контакты", "company_contacts"),
            ("Категории компаний", "company_category"),
            ("Документы", "company_documents"),
            ("Кошельки", "wallets"),
            ("Подписки", "subscriptions"),
            ("Доп. поля компаний", "company_attributes"),
            ("Раскрытые контакты", "contact_unlocks"),
            ("Отзывы", "reviews"),
        ]
        assert [(s.name, s.table) for s in LISTING_SHEETS] == [
            ("Объявления", "listings"),
            ("Фотографии", "listing_images"),
            ("Характеристики", "listing_attributes"),
            ("Статистика по дням", "listing_stats"),
            ("Избранное", "favorites"),
            ("Тендеры", "tenders"),
        ]
        assert all(s.columns and s.about for s in (*COMPANY_SHEETS, *LISTING_SHEETS))

    def test_заголовки_внутри_листа_не_повторяются(self):
        """Повторяющийся заголовок затёр бы соседний столбец."""
        for sheet in (*COMPANY_SHEETS, *LISTING_SHEETS):
            headers = [h for h, _, _ in sheet.columns]

            assert len(headers) == len(set(headers)), sheet.name


class TestЗапись:
    def test_книга_устроена_как_у_php_версии(self, tmp_path):
        sheet = Sheet(
            name="Компании",
            table="companies",
            about="Карточки",
            columns=(("ID", "id", "int"), ("Телефон", "phone", "text")),
        )
        book = Book(
            file=tmp_path / "b.xlsx", title="SAVDEX — компании", about="О книге", sheets=(sheet,)
        )

        write(book, [Collected(sheet=sheet, rows=[[1, "00998123"], [2, ""]])])

        workbook = load_workbook(book.file)

        assert workbook.sheetnames == ["Справка", "Компании"]

        page = workbook["Компании"]

        assert [c.value for c in page[1]] == ["ID", "Телефон"]
        assert page["B2"].value == "00998123"
        assert page["A2"].value == 1
        assert page.freeze_panes == "A2"
        assert page.auto_filter.ref == "A1:B3"
        assert page.column_dimensions["A"].width == 9
        assert page["A1"].font.b is True
        assert workbook["Справка"]["A1"].value == "SAVDEX — компании"


class TestСравнениеВыгрузок:
    """
    Сравнение книг PHP и Python — рабочий код, а не только тест.

    На боевом сервере каждая выгрузка из админки делается обеими
    версиями и сверяется этим сравнением.
    """

    @staticmethod
    def выгрузить(directory, rows):
        from datetime import UTC, datetime

        from savdex.export.workbooks import books

        directory.mkdir(parents=True, exist_ok=True)
        sheet = Sheet(name="Компании", table="companies", about="—", columns=(("ID", "id", "int"),))

        for book in books(directory, now=datetime(2026, 9, 26, 12, 0, tzinfo=UTC)):
            write(book, [Collected(sheet=sheet, rows=rows)])

    def test_одинаковые_выгрузки_совпадают(self, tmp_path):
        from savdex.export.compare import pair

        self.выгрузить(tmp_path / "php", [[1], [2]])
        self.выгрузить(tmp_path / "py", [[1], [2]])

        assert pair(tmp_path / "php", tmp_path / "py") == {"companies": [], "listings": []}

    def test_расхождение_названо_по_ячейке(self, tmp_path):
        from savdex.export.compare import pair

        self.выгрузить(tmp_path / "php", [[1], [2]])
        self.выгрузить(tmp_path / "py", [[1], [3]])

        result = pair(tmp_path / "php", tmp_path / "py")

        assert any("Компании!A3" in p for p in result["companies"])

    def test_пропавшая_книга_не_сходит_за_совпавшую(self, tmp_path):
        from savdex.export.compare import pair

        self.выгрузить(tmp_path / "php", [[1]])
        (tmp_path / "py").mkdir()

        result = pair(tmp_path / "php", tmp_path / "py")

        assert result["companies"] == ["нет книги companies в выгрузке Python"]
