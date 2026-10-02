# ruff: noqa: E501 — словари заголовков и образец перенесены из PHP как есть
"""
Загрузка товаров из книги Excel вместе с фотографиями — копия
App\\Services\\ListingWorkbookImport (с App\\Support\\WorkbookImages,
ListingWorkbookTemplate и частью ImportLanguage и CatalogLookup, которой
нет в savdex/tenders/importer.py), плюс ListingsTable::importWorkbooks —
несколько книг за раз.

Обычный импорт принимает только CSV, а в CSV картинку не положишь:
заказчик собирает каталог в Excel, вставляя фотографии прямо в ячейки.
Поэтому книга читается целиком — значения берёт openpyxl, фотографии
достаёт extract_photos из zip-архива книги, — и то и другое сходится
по номеру строки.

Книга — до пяти листов, по одному на язык. Лист узнаётся по имени
вкладки (SHEET_LOCALES); без узнаваемых имён главным считается первый
лист, как в книге на одном языке.

 — русский лист главный: с него берутся цена, валюта, категория,
   компания, город и фотографии — всё, что у товара одно на все языки;
 — остальные листы отдают только тексты: заголовок, описание, условия
   поставки и оплаты — строка к строке с русским листом. Лист, чьи
   строки не сходятся с русским, пропускается с заметкой (иначе переводы
   съедут на соседний товар), а товары загружаются без этого языка.

Пустое и неполное загрузку не останавливает: пустая ячейка
пропускается, строка, которую не загрузить (нет заголовка, нет
компании), идёт в отчёт, а остальные загружаются дальше; второй лист
на том же языке, лист перевода без шапки, «Номер», которого нет в базе,
и книга без листа «Русский» — заметка, а не отказ. Строки с названиями
столбцов нет или названия не узнаны — столбцы угадываются по содержимому
(_guess_columns), что чем стало — в отчёте.

Строка русского листа: «Номер» заполнен → правится это объявление;
иначе ищется по заголовку и компании; не нашлось → заводится новое
(без компании нельзя). Новое ждёт проверки (moderation), уже
опубликованное повторная загрузка не снимает. Фотографии добавляются
только объявлениям без фотографий; replace — заменить имеющиеся.

Отличия от PHP (осознанные):

- каждая строка — своя точка сохранения: сбой базы в одной строке не
  губит остальные и не оставляет полусохранённого объявления; счётчики
  строки учитываются только после успешной записи, а файлы прежних
  фотографий (replace) удаляются после неё — откат не оставит строк
  listing_images без файлов;
- XML частей книги с DOCTYPE не разбирается вовсе (у настоящих книг
  его нет) — вместо LIBXML_NONET без подстановки сущностей;
- книга, которую не открыть, — UnreadableWorkbookError (у Laravel —
  исключение openspout и ошибка в окне Filament).

Исправлено после отключения загрузки у Laravel (раздел Filament убран;
«Отложено до отключения Laravel» в docs/migration-to-python.md):
пропущенные ячейки (_Import.skipped) сбрасываются перед каждой строкой —
раньше заметка строки повторялась у всех следующих.

Добавлено сверх Laravel — для книг заявок, собранных с других площадок,
где нет продавца из справочника и столбца «Тип»:

- компания из окна загрузки (default_company), а без неё — служебная
  (Anjir Group, service_company_id) достаётся строкам, где компания не
  указана или не нашлась; существующему объявлению продавца не меняет;
  повторная загрузка с настоящей компанией передаёт ей заявку служебной,
  а не заводит вторую;
- тип новых объявлений, когда ячейки «Тип» нет: по заголовку («Куплю…»,
  «Требуется…» — запрос, «Продам…» — предложение), иначе из окна
  (default_type); ячейка «Тип» понимает и «Запрос на закупку»,
  «Покупаем», а непонятное слово — заметка, а не «Предложение» молча;
- «Раздел → Подраздел»: «Прочее» — это «Другое» каталога, неизвестный
  подраздел — сам раздел с заметкой (listing_category);
- длинный заголовок ищется обрезанным, как сохранён: повторная
  загрузка не заводит дубль.

Запись — в чужие таблицы listings и listing_images (SHARED_WRITES),
журнал — как AuditObserver у администратора и строка «imported», как
действие importWorkbook в ListingsTable.

Сверка с тестом Laravel (tests/Feature/Admin/ListingWorkbookImportTest.php) —
tests/test_listing_workbook.py.
"""

from __future__ import annotations

import io
import json
import logging
import posixpath
import re
import tempfile
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, TypedDict
from xml.etree import ElementTree

from django.db import connection, transaction

from savdex import access, audit
from savdex.guards import allowed_writes
from savdex.tenders import importer
from savdex.tenders.slug import slugify
from savdex.web import image_store, locales
from savdex.web.cabinet import _rows
from savdex.web.eloquent import _cast_same, _written
from savdex.web.listing_actions import LIFETIME_DAYS, _now, _search_text, _stamp
from savdex.web.listing_image_actions import MAX_IMAGES
from savdex.web.seo import _width
from savdex.web.validation import php_float
from savdex.web.wizard_actions import CASTS, CURRENCIES

log = logging.getLogger(__name__)

# ── Словари (ImportLanguage) ─────────────────────────────────────────

#: ImportLanguage::LISTING_HEADERS — как называют столбец в таблице
#: товаров; списки уже нормализованы
LISTING_HEADERS: dict[str, list[str]] = {
    "id": ["номер", "номер объявления", "id", "raqam", "no", "编号"],
    "title": ["заголовок", "название", "наименование", "название товара", "наименование товара",
        "товар", "продукция", "позиция",
        "title", "name", "product", "product name", "item",
        "nomi", "mahsulot", "mahsulot nomi", "başlık", "ürün", "ürün adı", "名称", "产品"],
    "company": ["компания", "поставщик", "продавец", "организация", "фирма", "инн",
        "company", "supplier", "seller", "vendor", "organization",
        "kompaniya", "yetkazib beruvchi", "sotuvchi", "şirket", "firma", "tedarikçi", "公司", "供应商"],
    "category_id": ["категория", "категории", "категория товара", "раздел", "рубрика", "группа", "отрасль",
        "category", "categories", "section", "group", "industry",
        "kategoriya", "turkum", "bo'lim", "soha", "kategori", "bölüm", "类别", "分类"],
    "type": ["тип", "вид", "тип объявления", "предложение или запрос",
        "type", "kind", "listing type", "turi", "tür", "ilan türü", "类型"],
    "price": ["цена", "стоимость", "цена за единицу", "price", "cost", "unit price",
        "narx", "narxi", "fiyat", "birim fiyat", "价格"],
    "currency": ["валюта", "currency", "valyuta", "pul birligi", "para birimi", "货币"],
    "unit": ["единица", "единица измерения", "ед. изм.", "ед изм", "ед",
        "unit", "measure", "uom", "birlik", "o'lchov birligi", "birim", "单位"],
    "min_order": ["минимальный заказ", "мин. заказ", "мин заказ", "минимальная партия",
        "minimum order", "min order", "moq",
        "eng kam buyurtma", "minimal buyurtma", "minimum sipariş", "最小起订量"],
    "city_id": ["город", "регион", "область", "city", "region", "shahar", "viloyat", "şehir", "城市"],
    "description": ["описание", "подробности", "характеристики", "условия",
        "description", "details", "specs", "specification",
        "tavsif", "izoh", "xususiyatlari", "açıklama", "özellikler", "描述", "说明"],
    "delivery_terms": ["условия поставки", "условия доставки", "поставка", "доставка", "отгрузка",
        "delivery terms", "delivery", "shipping", "shipping terms",
        "yetkazib berish shartlari", "yetkazib berish", "teslimat koşulları", "teslimat", "teslimat şartları",
        "交货条件", "交付条件", "运输条件", "交货"],
    "payment_terms": ["условия оплаты", "оплата", "порядок оплаты",
        "payment terms", "payment",
        "to'lov shartlari", "to'lov", "ödeme koşulları", "ödeme", "ödeme şartları",
        "付款条件", "支付条件", "付款"],
}  # fmt: skip

#: ImportLanguage::SHEET_LOCALES — как называют лист книги с текстами на
#: одном языке; лист узнаётся по имени вкладки, а не по порядку
SHEET_LOCALES: dict[str, list[str]] = {
    "ru": ["ru", "rus", "ru-ru", "русский", "рус", "russian", "ruscha", "rusça", "rusca", "俄语", "俄文"],
    "en": ["en", "eng", "en-us", "en-gb", "english", "английский", "англ", "inglizcha", "ingliz", "ingilizce", "英语", "英文"],
    "uz": ["uz", "uzb", "uz-uz", "o'zbekcha", "o'zbek", "uzbek", "uzbekcha", "ozbekcha", "узбекский", "узбек", "özbekçe", "ozbekce", "乌兹别克语"],
    "zh": ["zh", "zh-cn", "zh-hans", "cn", "chinese", "китайский", "кит", "xitoycha", "xitoy", "çince", "cince", "中文", "汉语", "简体中文"],
    "tr": ["tr", "tur", "tr-tr", "türkçe", "turkce", "turkish", "турецкий", "тур", "turkcha", "turk", "土耳其语", "土耳其文"],
}  # fmt: skip

#: ListingWorkbookImport::type — слова запроса; всё остальное — предложение
_DEMAND = frozenset({
    "запрос", "спрос", "закупка", "покупка", "куплю", "потребность",
    "demand", "request", "rfq", "buy", "buying", "purchase",
    "talab", "sotib olish", "xarid", "talep", "alım", "satın alma",
    "采购", "求购", "需求",
})  # fmt: skip

# ── Объявление (App\Models\Listing) ─────────────────────────────────

#: Listing::TRANSLATABLE — поля с версией на каждом языке (колонка *_i18n)
TRANSLATABLE = ("title", "description", "delivery_terms", "payment_terms")

#: Listing::MAX_LENGTH — предел длины, один на кабинет, админку и загрузку
MAX_LENGTH = {"title": 90, "description": 5000, "delivery_terms": 2000, "payment_terms": 2000}

#: Как поле называется в заметке об обрезке (ListingWorkbookImport::fits)
_LABELS = {
    "title": "заголовок",
    "description": "описание",
    "delivery_terms": "условия поставки",
    "payment_terms": "условия оплаты",
}

# ── Образец (ListingWorkbookTemplate) ───────────────────────────────

#: ListingWorkbookTemplate::HEADERS — столбцы русского листа
HEADERS = [
    "Номер", "Название", "Компания", "Категория", "Тип", "Цена", "Валюта",
    "Единица", "Минимальный заказ", "Город", "Описание",
    "Условия поставки", "Условия оплаты", "Фото",
]  # fmt: skip

_EXAMPLE = [
    "", "Кирпич керамический М150", "ООО «Стройбаза»", "Стройматериалы → Кирпич и блоки",
    "Предложение", "1200", "UZS", "шт", "5000", "Ташкент",
    "Полнотелый, марка М150, отгрузка с завода.",
    "Самовывоз со склада, доставка по Ташкентской области.",
    "Предоплата 50 %, остаток по факту отгрузки.", "",
]  # fmt: skip

#: Языковые листы образца: столбцы и пример — на своём языке
_TRANSLATIONS: dict[str, tuple[list[str], list[str]]] = {
    "en": (
        ["Title", "Description", "Delivery terms", "Payment terms"],
        ["Ceramic brick M150", "Solid, grade M150, shipped from the plant.",
            "Pick-up from the warehouse, delivery across the Tashkent region.",
            "50 % prepayment, the rest on shipment."],
    ),
    "uz": (
        ["Nomi", "Tavsif", "Yetkazib berish shartlari", "To'lov shartlari"],
        ["Keramik g‘isht M150", "To‘liq, M150 markasi, zavoddan jo‘natish.",
            "Ombordan o‘zi olib ketish, Toshkent viloyati bo‘ylab yetkazish.",
            "50 % oldindan to‘lov, qolgani jo‘natish bo‘yicha."],
    ),
    "zh": (
        ["名称", "描述", "交货条件", "付款条件"],
        ["M150 陶瓷砖", "实心砖，M150 标号，工厂直发。", "仓库自提，塔什干州范围内配送。", "预付 50 %，余款发货后结清。"],
    ),
    "tr": (
        ["Başlık", "Açıklama", "Teslimat koşulları", "Ödeme koşulları"],
        ["Seramik tuğla M150", "Dolu, M150 sınıfı, fabrikadan sevkiyat.",
            "Depodan teslim, Taşkent bölgesine dağıtım.", "%50 peşin, kalanı sevkiyatta."],
    ),
}  # fmt: skip

#: Имя файла образца, как у кнопки «Скачать образец»
TEMPLATE_NAME = "savdex-tovary-obrazec.xlsx"

#: ListingsTable::MAX_WORKBOOKS — сколько книг за раз
MAX_WORKBOOKS = 10


def sheet_name(locale: str) -> str:
    """ListingWorkbookTemplate::sheetName: имя вкладки — как в переключателе языка."""
    return locales.ALL[locale]["label"]


def template_bytes() -> bytes:
    """
    ListingWorkbookTemplate::write: образец книги — русский лист с ценами
    и фотографиями, за ним по листу на язык, строка к строке.
    """
    from openpyxl import Workbook

    book = Workbook()
    page: Any = book.active
    page.title = sheet_name("ru")

    def add(target: Any, values: list[str]) -> None:  # noqa: ANN401
        # Пустая ячейка — пустая, а не строка нулевой длины (EmptyCell)
        target.append([value if value != "" else None for value in values])

    add(page, HEADERS)
    add(page, _EXAMPLE)

    for locale, (headers, example) in _TRANSLATIONS.items():
        sheet = book.create_sheet(sheet_name(locale))
        add(sheet, headers)
        add(sheet, example)

    out = io.BytesIO()
    book.save(out)

    return out.getvalue()


# ── Итог ─────────────────────────────────────────────────────────────


class WorkbookResult(TypedDict):
    """Итог загрузки: строки, новые, обновлённые, фотографии, ошибки, заметки."""

    rows: int
    created: int
    updated: int
    photos: int
    errors: list[str]
    notes: list[str]


class UnreadableWorkbookError(ValueError):
    """Файл не открывается как книга Excel (.xlsx)."""


def _empty() -> WorkbookResult:
    return {"rows": 0, "created": 0, "updated": 0, "photos": 0, "errors": [], "notes": []}


def import_workbooks(
    files: Sequence[tuple[str, bytes]],
    *,
    admin_id: int,
    replace: bool,
    ip: str | None = None,
    default_company: int | None = None,
    default_type: str = "supply",
    publish: bool = False,
) -> WorkbookResult:
    """
    Действие importWorkbook в ListingsTable: книги подряд, итоги
    складываются; строка журнала «imported» с числом книг и итогами.

    files — (имя файла, содержимое). Когда книг несколько, каждая строка
    ошибок и заметок начинается с имени файла. Пустой список — ничего не
    делается (у Filament — «Файл не получен»). Книга, которую не открыть, —
    UnreadableWorkbookError; загруженное из книг до неё остаётся.
    default_company и default_type — из окна загрузки: компания строк
    без компании (или с компанией не из справочника) и тип новых
    объявлений, когда в книге нет столбца «Тип».
    """
    total = _empty()

    if not files:
        return total

    actor = _actor(admin_id)
    many = len(files) > 1

    for name, content in files:
        result = import_workbook(
            content,
            author_id=admin_id,
            replace=replace,
            actor=actor,
            ip=ip,
            default_company=default_company,
            default_type=default_type,
            publish=publish,
        )

        total["rows"] += result["rows"]
        total["created"] += result["created"]
        total["updated"] += result["updated"]
        total["photos"] += result["photos"]

        # Из какой книги строка — понятно только когда их несколько
        prefix = f"{name}, " if many else ""
        total["errors"].extend(prefix + (_lcfirst(e) if many else e) for e in result["errors"])
        total["notes"].extend(prefix + (_lcfirst(n) if many else n) for n in result["notes"])

    # След в журнале: загрузка создаёт записи пачкой, минуя формы
    audit.record(
        connection,
        action="imported",
        section="listings",
        actor=actor,
        note=(
            f"Книг: {len(files)}, создано: {total['created']}, обновлено: {total['updated']}, "
            f"фотографий: {total['photos']}"
        ),
        ip=ip,
    )

    return total


def import_workbook(
    content: bytes,
    *,
    author_id: int,
    replace: bool = False,
    actor: access.Admin | None = None,
    ip: str | None = None,
    default_company: int | None = None,
    default_type: str = "supply",
    publish: bool = False,
) -> WorkbookResult:
    """
    ListingWorkbookImport::run — одна книга. author_id — автор новых
    объявлений; actor — кто загружает: у администратора правки пишутся
    в журнал (AuditObserver), без него — нет.
    """
    return _Import(
        author_id=author_id,
        replace=replace,
        observer=actor if actor is not None and actor.is_admin else None,
        ip=ip,
        default_company=default_company,
        default_type="demand" if default_type == "demand" else "supply",
        publish=publish,
    ).run(content)


def _actor(user_id: int) -> access.Admin | None:
    """Auth::user() действия: автор записи журнала."""
    found = _rows(
        "select id, name, email, is_admin, admin_role, status from users where id = %s",
        [user_id],
    )

    if not found:
        return None

    user = found[0]

    return access.Admin(
        id=int(user["id"]),
        name=str(user["name"]),
        email=str(user["email"]),
        is_admin=bool(user["is_admin"]),
        role=user["admin_role"],
        status=str(user["status"]),
    )


# ── Мелочи PHP ───────────────────────────────────────────────────────

_PHP_TRIM = " \t\n\r\0\x0b"


def _trim(value: str) -> str:
    """trim() у PHP: только ASCII-пробелы, неразрывный остаётся."""
    return value.strip(_PHP_TRIM)


def _lcfirst(value: str) -> str:
    """lcfirst() у PHP: только латинская буква — кириллицу он не трогает."""
    return value[0].lower() + value[1:] if value[:1].isascii() and value[:1].isupper() else value


def _strimwidth(value: str, width: int, marker: str = "…") -> str:
    """mb_strimwidth($value, 0, $width, $marker): по ширине знаков, с отметкой."""
    if sum(_width(char) for char in value) <= width:
        return value

    room = width - sum(_width(char) for char in marker)
    kept: list[str] = []
    used = 0

    for char in value:
        used += _width(char)

        if used > room:
            break

        kept.append(char)

    return "".join(kept) + marker


def _php_int(value: str | None) -> int:
    """(int) строки у PHP: число в начале строки, иначе ноль."""
    found = re.match(r"[ \t\n\r\v\f]*([+-]?[0-9]+)", value or "")

    return int(found.group(1)) if found else 0


def _text(value: Any, epoch: datetime) -> str:  # noqa: ANN401
    """
    ListingWorkbookImport::text — значение ячейки строкой.

    Excel отдаёт целые числа дробными: «Номер 12» иначе превратился бы
    в «12.0» и не нашёл объявление. Дата — «d.m.Y»; время без даты
    openspout отдаёт днём начала отсчёта книги — так и здесь.
    """
    if isinstance(value, date):
        return f"{value.day:02d}.{value.month:02d}.{value.year:04d}"

    if isinstance(value, time):
        return f"{epoch.day:02d}.{epoch.month:02d}.{epoch.year:04d}"

    if isinstance(value, bool):
        return "да" if value else "нет"

    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else _trim(php_float(value))

    if isinstance(value, timedelta):
        return _trim(str(value))

    return _trim(str(value if value is not None else ""))


# ── Справочники (CatalogLookup: город и компания) ───────────────────


def _without_city_prefix(value: str) -> str:
    """«г. ташкент», «ташкент г.», «toshkent sh.» → «ташкент»."""
    value = re.sub(r"^(?:г|гор|город|sh|shahri|shahar)\.?\s+", "", value)
    value = re.sub(r"\s+(?:г|гор|город|sh|shahri|shahar)\.?$", "", value)

    return value.strip()


def city_id(value: str | None) -> int | None:
    """CatalogLookup::cityId: город по slug или названию на любом языке."""
    needle = _without_city_prefix(importer.normalize(value))

    if needle == "":
        return None

    by_slug = _rows("select id from cities where slug = %s limit 1", [needle])

    if by_slug:
        return int(by_slug[0]["id"])

    for row in _rows("select city_id, name from city_translations order by id"):
        if _without_city_prefix(importer.normalize(row["name"])) == needle:
            return int(row["city_id"])

    return None


def service_company_id() -> int | None:
    """Служебная компания — как у витрины (savdex/web/platform.py)."""
    from savdex.web.platform import service_company_id as lookup

    return lookup()


def company_name(company: int | None) -> str:
    rows = _rows("select name from companies where id = %s", [company]) if company else []

    return str(rows[0]["name"]) if rows else ""


def company_id(value: str | None) -> int | None:
    """
    CatalogLookup::companyId: по ИНН, затем по названию или юридическому
    названию; компании в корзине не ищутся (SoftDeletes).
    """
    raw = _trim(str(value or ""))
    needle = importer.normalize(raw)

    if needle == "":
        return None

    if re.fullmatch(r"[0-9]{6,20}", raw):
        by_tin = _rows(
            "select id from companies where tin = %s and deleted_at is null limit 1", [raw]
        )

        if by_tin:
            return int(by_tin[0]["id"])

    for row in _rows(
        "select id, name, legal_name, tin from companies where deleted_at is null order by id"
    ):
        if (
            importer.normalize(row["name"]) == needle
            or importer.normalize(str(row["legal_name"] or "")) == needle
            or str(row["tin"] or "") == raw
        ):
            return int(row["id"])

    return None


#: Подраздел «всё остальное» называют по-разному: в каталоге «Другое»,
#: в книгах — «Прочее», «Разное»
_OTHER = ("другое", "прочее", "разное", "иное", "остальное", "other", "boshqa", "diğer", "其他")


def _category_names(category: int) -> set[str]:
    rows = _rows(
        "select c.slug, t.name from categories c "
        "left join category_translations t on t.category_id = c.id where c.id = %s",
        [category],
    )

    return {importer.normalize(r["slug"]) for r in rows} | {
        importer.normalize(r["name"]) for r in rows if r["name"]
    }


def _cut(value: str, limit: int) -> str:
    """Обрезка по границе слова: хвост теряется, а не смысл."""
    if len(value) <= limit:
        return value

    cut = value[:limit]
    space = cut.rfind(" ")

    return (cut[:space] if space > limit / 2 else cut).rstrip(" ,.;-")


def listing_category(value: str | None) -> tuple[int | None, str | None]:
    """
    Категория строки и заметка, если она подобрана не буквально.

    «Раздел → Подраздел» ищется внутри найденного раздела: подраздел
    «Прочее» — это «Другое» каталога, а подраздела, которого нет, —
    сам раздел (объявление не теряет категорию из-за одного слова).
    Без раздела — как раньше: по названию на любом языке.
    """
    raw = _trim(str(value or ""))
    path = [p for p in re.split(r"\s*(?:→|->|>|/|\\|\|)\s*", raw) if _trim(p) != ""]

    if len(path) < 2:
        return importer.category_id(raw), None

    parent = None

    # Раздел — сперва среди верхних: «Другое» есть и разделом, и
    # подразделом почти в каждом разделе
    if len(path) == 2:
        needle = importer.normalize(path[0])
        parent = next(
            (
                int(top["id"])
                for top in _rows(
                    "select id from categories where parent_id is null order by sort, id"
                )
                if needle in _category_names(int(top["id"]))
            ),
            None,
        )

    parent = parent or importer.category_id(" → ".join(path[:-1]))

    if parent is None:
        return importer.category_id(raw), None

    leaf = importer.normalize(path[-1])
    wanted = {leaf} | (set(_OTHER) if leaf in _OTHER else set())

    for child in _rows(
        "select id from categories where parent_id = %s order by sort, id", [parent]
    ):
        if _category_names(int(child["id"])) & wanted:
            return int(child["id"]), None

    names = _rows(
        "select name from category_translations where category_id = %s and locale = 'ru'", [parent]
    )
    section = str(names[0]["name"]) if names else _trim(path[-2])

    return parent, (
        f"подраздел «{_trim(path[-1])}» не найден в каталоге — объявление в разделе «{section}»"
    )


def sheet_locale(name: str) -> str | None:
    """ImportLanguage::sheetLocale: язык листа по имени вкладки; None — не про язык."""
    needle = importer.normalize(name)

    return next((code for code, aliases in SHEET_LOCALES.items() if needle in aliases), None)


# ── Фотографии в книге (WorkbookImages) ─────────────────────────────

_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
_DOC_RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_DRAWING = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_MAIN = "http://schemas.openxmlformats.org/drawingml/2006/main"
_SHEET = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def extract_photos(content: bytes, into: Path, sheet: str | None = None) -> dict[int, list[Path]]:
    """
    WorkbookImages::extract: картинки одного листа — в каталог, по строкам.

    xlsx — zip с XML: xl/workbook.xml (листы), xl/worksheets/sheetN.xml
    и его _rels (ссылка на рисунки), xl/drawings/drawingN.xml (где стоит
    картинка), xl/media/* (сами файлы). Картинка принадлежит строке, где
    её верхний левый угол; absoluteAnchor (без ячейки) пропускается,
    внешние картинки по ссылке — тоже. Лист — по имени вкладки, без
    имени — первый. Файлы кладутся на диск, а не держатся в памяти.

    Номер строки — с единицы, как в Excel; в строке — по порядку колонок.
    """
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except (zipfile.BadZipFile, OSError):
        return {}

    with archive:
        part = _sheet_path(archive, sheet)

        if part is None:
            return {}

        drawing = _related_part(archive, part, "drawing")

        if drawing is None:
            return {}

        root = _xml(archive, drawing)

        if root is None:
            return {}

        media = _relation_targets(archive, drawing)
        rows: dict[int, list[tuple[int, str]]] = {}

        for anchor in _children(root, _DRAWING):
            placed = _anchored(anchor, media)

            if placed is not None:
                row, column, target = placed
                rows.setdefault(row, []).append((column, target))

        return _save_photos(archive, rows, into)


def _anchored(anchor: ElementTree.Element, media: dict[str, str]) -> tuple[int, int, str] | None:
    """Строка, колонка и файл картинки одного якоря."""
    corner = anchor.find(f"{{{_DRAWING}}}from")
    picture = anchor.find(f"{{{_DRAWING}}}pic")

    if corner is None or picture is None:
        return None

    fill = picture.find(f"{{{_DRAWING}}}blipFill")
    blip = fill.find(f"{{{_MAIN}}}blip") if fill is not None else None

    if blip is None:
        return None

    target = media.get(blip.get(f"{{{_DOC_RELS}}}embed", ""))

    if target is None:
        return None

    # В файле строки и колонки считаются от нуля
    return (
        _php_int(corner.findtext(f"{{{_DRAWING}}}row")) + 1,
        _php_int(corner.findtext(f"{{{_DRAWING}}}col")),
        target,
    )


def _save_photos(
    archive: zipfile.ZipFile, rows: dict[int, list[tuple[int, str]]], into: Path
) -> dict[int, list[Path]]:
    """Достать файлы из книги и разложить по строкам; первая в строке — обложка."""
    try:
        into.mkdir(parents=True, exist_ok=True)
    except OSError:
        return {}

    saved: dict[int, list[Path]] = {}
    number = 0

    for row in sorted(rows):
        for _column, part in sorted(rows[row], key=lambda picture: picture[0]):
            try:
                binary = archive.read(part)
            except Exception:  # как getFromName(): битая или отсутствующая часть — false
                continue

            if binary == b"":
                continue

            number += 1
            path = into / f"{number}-{posixpath.basename(part)}"

            try:
                path.write_bytes(binary)
            except OSError:
                continue

            saved.setdefault(row, []).append(path)

    return saved


def _sheet_path(archive: zipfile.ZipFile, name: str | None) -> str | None:
    """Путь к XML листа: по имени вкладки (буквально), без имени — первого."""
    root = _xml(archive, "xl/workbook.xml")

    if root is None:
        return None

    sheets = root.find(f"{{{_SHEET}}}sheets")

    if sheets is None:
        return None

    chosen = next(
        (s for s in sheets.findall(f"{{{_SHEET}}}sheet") if name is None or s.get("name") == name),
        None,
    )

    if chosen is None:
        return None

    return _relation_targets(archive, "xl/workbook.xml").get(chosen.get(f"{{{_DOC_RELS}}}id", ""))


def _related_part(archive: zipfile.ZipFile, part: str, kind: str) -> str | None:
    """Часть книги, на которую ссылается другая часть: рисунки листа."""
    rels = _xml(archive, _rels_path(part))

    if rels is None:
        return None

    for relation in _children(rels, _RELS):
        if relation.get("Type", "").endswith("/" + kind):
            return _resolve(part, relation.get("Target", ""))

    return None


def _relation_targets(archive: zipfile.ZipFile, part: str) -> dict[str, str]:
    """Идентификатор связи → путь к файлу внутри книги; внешние ссылки — мимо."""
    rels = _xml(archive, _rels_path(part))

    if rels is None:
        return {}

    return {
        relation.get("Id", ""): _resolve(part, relation.get("Target", ""))
        for relation in _children(rels, _RELS)
        # Ходить за картинками в интернет загрузка не должна
        if relation.get("TargetMode", "") != "External"
    }


def _dirname(part: str) -> str:
    """dirname() у PHP: без каталога — «.»."""
    return posixpath.dirname(part) or "."


def _rels_path(part: str) -> str:
    directory = _dirname(part)
    directory = "" if directory == "." else directory + "/"

    return f"{directory}_rels/{posixpath.basename(part)}.rels"


def _resolve(part: str, target: str) -> str:
    """Ссылка внутри книги — относительно той части, где написана."""
    if target.startswith("/"):
        return target.lstrip("/")

    segments: list[str] = []

    for segment in f"{_dirname(part)}/{target}".split("/"):
        if segment in ("", "."):
            continue

        if segment == "..":
            if segments:
                segments.pop()

            continue

        segments.append(segment)

    return "/".join(segments)


def _children(element: ElementTree.Element, namespace: str) -> list[ElementTree.Element]:
    """->children($ns) у SimpleXML: прямые потомки из пространства имён."""
    return [child for child in element if child.tag.startswith(f"{{{namespace}}}")]


def _xml(archive: zipfile.ZipFile, part: str) -> ElementTree.Element | None:
    """
    Часть книги как XML; нет, пусто или не разобралось — None.

    Файл чужой: с DOCTYPE не разбирается вовсе — подстановка сущностей
    превращает «безобидный» XML в бомбу из вложенных сущностей, а у
    настоящих книг DOCTYPE не бывает.
    """
    if part == "":
        return None

    try:
        source = archive.read(part)
    except Exception:
        return None

    if source.strip(_PHP_TRIM.encode()) == b"" or b"<!DOCTYPE" in source or b"<!ENTITY" in source:
        return None

    try:
        return ElementTree.fromstring(source)
    except Exception:  # как @simplexml_load_string: любой сбой разбора — false
        return None


# ── Листы книги ──────────────────────────────────────────────────────


@dataclass
class _Row:
    number: int
    fields: dict[str, str]


@dataclass
class _Sheet:
    name: str
    locale: str | None
    header: int | None = None
    last: int = 0
    rows: dict[int, _Row] = field(default_factory=dict)
    #: Что лежит над таблицей без шапки: (номер строки, ячейки)
    sample: tuple[int, list[str]] | None = None
    #: Строки листа, пока шапка не нашлась, — текстом (номер строки,
    #: ячейки): без шапки по ним столбцы угадываются по содержимому
    raw: list[tuple[int, list[str]]] = field(default_factory=list)


def _columns(values: Iterable[Any], epoch: datetime) -> dict[int, str]:
    """Шапка таблицы: номер колонки → поле объявления, каждое поле — раз."""
    columns: dict[int, str] = {}
    taken: set[str] = set()

    for index, value in enumerate(values):
        header = _text(value, epoch)

        for name, aliases in LISTING_HEADERS.items():
            if name in taken or not importer.matches(header, aliases):
                continue

            columns[index] = name
            taken.add(name)

            break

    return columns


def _read_sheet(sheet: _Sheet, rows: Iterable[list[Any]], epoch: datetime) -> None:
    """
    ListingWorkbookImport::readSheet: шапка и всё, что под ней.

    Номер строки — как в Excel (по нему сходятся фотографии), смещение
    от шапки — ключ (по нему сходятся листы). Пока шапки нет, запоминается
    образец того, что над ней: первая строка хотя бы с двумя ячейками,
    а без таких — первая непустая.
    """
    columns: dict[int, str] | None = None
    sample: tuple[int, list[str]] | None = None
    number = 0

    for values in rows:
        number += 1

        if columns is None:
            found = _columns(values, epoch)
            texts = [_text(v, epoch) for v in values]
            cells = [text for text in texts if text != ""]

            # Шапка — строка, где узнаны хотя бы два столбца или все её
            # ячейки. Одно совпадение в строке данных («Доставка, варианты»
            # похоже на столбец «Доставка») шапкой её не делает
            if found and len(found) < 2 and len(found) < len(cells):
                found = {}

            if not found:
                sheet.raw.append((number, texts))

                if cells and (sample is None or (len(sample[1]) < 2 and len(cells) >= 2)):
                    sample = (number, cells)
            else:
                columns = found
                sheet.header = number

            continue

        assert sheet.header is not None
        fields: dict[str, str] = {}

        for index, name in columns.items():
            text = _text(values[index] if index < len(values) else None, epoch)

            if text != "":
                fields[name] = text

        offset = number - sheet.header
        sheet.rows[offset] = _Row(number, fields)

        if fields:
            sheet.last = offset

    sheet.sample = sample if sheet.header is None else None


#: Приметы столбцов без шапки: условия поставки и оплаты различаются по словам
_DELIVERY_WORDS = (
    "поставк",
    "доставк",
    "самовывоз",
    "отгрузк",
    "exw",
    "fca",
    "fob",
    "cif",
    "cfr",
    "cpt",
    "cip",
    "dap",
    "dpu",
    "ddp",
    "incoterms",
    "склад",
    "порт",
    "delivery",
    "shipping",
)
_PAYMENT_WORDS = (
    "оплат",
    "предоплат",
    "постоплат",
    "аккредитив",
    "l/c",
    "перечислен",
    "наличн",
    "рассрочк",
    "безнал",
    "payment",
    "prepayment",
    "letter of credit",
    "tt ",
    "t/t",
)
#: Организационная форма в названии — примета столбца компании
_LEGAL_FORM = re.compile(
    r"(?<![\w])(ООО|ОАО|ЗАО|ПАО|АО|ТОО|ИП|ЧП|МЧЖ|ХК|LLC|LLP|Ltd|GmbH|Inc|Corp|JSC|OOO)(?![\w])",
    re.IGNORECASE,
)
_SUPPLY = frozenset({"предложение", "продаю", "продажа", "supply", "offer", "sell", "sotaman"})

#: Начало слов «покупаю» в ячейке «Тип»: ячейку пишут как придётся —
#: «Запрос на закупку», «Покупаем», «Спрос», «Buy request», и точное
#: совпадение со списком _DEMAND отдавало всё это в «Предложение»
_DEMAND_STEMS = (
    "запрос", "спрос", "закуп", "покуп", "купл", "куп", "потребн", "приобрет",
    "demand", "request", "rfq", "buy", "purchas", "wanted",
    "talab", "sotib", "xarid", "talep", "alım", "satın",
    "采购", "求购", "需求",
)  # fmt: skip
_SUPPLY_STEMS = (
    "предлож", "прода", "реализ", "supply", "offer", "sell", "sale",
    "sotaman", "sotuv", "satış", "satılık", "出售", "供应",
)  # fmt: skip

#: Заголовок заявки на покупку: «Куплю…», «Требуется поставщик…», «Ищем…».
#: По нему тип ставится, когда в книге нет столбца «Тип» — так собраны
#: книги заявок с других площадок
_DEMAND_TITLE = re.compile(
    r"^\W*(куплю|купим|покупаем|покупаю|закупаем|закупим|закупка|приобрету|приобретём|"
    r"приобретем|приобретаем|требуется|требуются|нужен|нужна|нужно|нужны|ищу|ищем|"
    r"запрос|wanted|buy|buying|looking for|need|needed|rfq|request|"
    r"sotib olaman|sotib olamiz|xarid qilamiz|kerak|alıyoruz|alınır|aranıyor|求购|采购)(?![\w])",
    re.IGNORECASE,
)
_SUPPLY_TITLE = re.compile(
    r"^\W*(продам|продаю|продаём|продаем|продаётся|продается|продаются|предлагаем|предлагаю|"
    r"реализуем|в наличии|sell|selling|for sale|offer|sotaman|sotamiz|sotiladi|satılık|"
    r"satıyoruz|出售|供应)(?![\w])",
    re.IGNORECASE,
)


def type_from_cell(value: str) -> str | None:
    """
    Тип из ячейки «Тип»: «demand», «supply» или None, если слово не
    распознано. Сначала точные слова, затем начало слова — «Запрос на
    закупку», «Покупаем», «Продажа со склада».
    """
    text = importer.normalize(value)

    if text in _DEMAND:
        return "demand"

    if text in _SUPPLY:
        return "supply"

    words = re.findall(r"\w+", text)

    if any(word.startswith(_DEMAND_STEMS) for word in words) or any(
        stem in text for stem in ("采购", "求购", "需求")
    ):
        return "demand"

    if any(word.startswith(_SUPPLY_STEMS) for word in words) or any(
        stem in text for stem in ("出售", "供应")
    ):
        return "supply"

    return None


def type_from_title(title: str) -> str | None:
    """Тип по началу заголовка: «Куплю…» — demand, «Продам…» — supply, иначе None."""
    if _DEMAND_TITLE.search(title):
        return "demand"

    if _SUPPLY_TITLE.search(title):
        return "supply"

    return None


#: Как назвать столбец в отчёте
_GUESSED = {
    "title": "заголовок",
    "description": "описание",
    "category_id": "категория",
    "company": "компания",
    "city_id": "город",
    "type": "тип",
    "price": "цена",
    "currency": "валюта",
    "delivery_terms": "условия поставки",
    "payment_terms": "условия оплаты",
}


def _letter(index: int) -> str:
    """Буква столбца Excel: 0 → A, 26 → AA."""
    letters = ""
    index += 1

    while index:
        index, rest = divmod(index - 1, 26)
        letters = chr(65 + rest) + letters

    return letters


def _share(values: list[str], test: Any) -> float:  # noqa: ANN401
    return sum(1 for v in values if test(v)) / len(values) if values else 0.0


def _guess_columns(sheet: _Sheet) -> str | None:
    """
    Лист без строки с названиями столбцов: столбцы угадываются по
    содержимому, чтобы книга не отклонялась из-за потерянной шапки.

    Данные — строки хотя бы с двумя заполненными ячейками, начиная с
    первой такой (заголовок отчёта в одной ячейке над таблицей не в счёт).
    Столбец узнаётся по приметам: «Раздел → Подраздел» или название из
    каталога — категория; код валюты — валюта; число — цена; город или
    компания из справочника; «куплю / продаю» — тип; самый длинный текст —
    описание; первый слева текст — заголовок; слова «CIF, доставка,
    самовывоз» и «оплата, аккредитив» — условия поставки и оплаты.
    Без заголовка — не угадано (None). Иначе строки листа заполняются
    и возвращается заметка для отчёта: что каким столбцом стало.
    """
    body = [(n, cells) for n, cells in sheet.raw if sum(1 for c in cells if c != "") >= 2]

    if not body:
        return None

    # Первая строка — непризнанная шапка («Наименование позиции»,
    # «Сколько стоит»), если под её коротким словом в каком-то столбце
    # стоят числа или длинные тексты: тогда это не товар, а названия
    skipped_header: int | None = None

    if len(body) >= 2 and _looks_like_header(body[0][1], [cells for _, cells in body[1:]]):
        skipped_header = body[0][0]
        body = body[1:]

    first = body[0][0]
    width = max(len(cells) for _, cells in body)
    rows = [(n, cells) for n, cells in sheet.raw if n >= first]
    values = {
        i: [cells[i] for _, cells in body if i < len(cells) and cells[i] != ""][:40]
        for i in range(width)
    }
    filled = {i: v for i, v in values.items() if v}
    columns: dict[int, str] = {}

    def free() -> list[int]:
        return [i for i in filled if i not in columns]

    def number(v: str) -> bool:
        return re.fullmatch(r"[0-9][0-9\s.,]*", v.strip()) is not None

    def average(i: int) -> float:
        return sum(len(v) for v in filled[i]) / len(filled[i])

    checks: list[tuple[str, Any]] = [
        ("category_id", lambda v: "→" in v or "->" in v or listing_category(v)[0] is not None),
        ("currency", lambda v: len(v) <= 6 and importer.currency(v) in CURRENCIES),
        (
            "type",
            lambda v: importer.normalize(v) in _DEMAND or importer.normalize(v) in _SUPPLY,
        ),
        (
            "company",
            lambda v: (
                len(v) <= 190 and (_LEGAL_FORM.search(v) is not None or company_id(v) is not None)
            ),
        ),
        ("city_id", lambda v: len(v) <= 80 and city_id(v) is not None),
    ]

    for name, test in checks:
        best = max(free(), key=lambda i: _share(filled[i], test), default=None)

        if best is not None and _share(filled[best], test) >= 0.6:
            columns[best] = name

    # Числа: первый слева столбец сплошь из целых — скорее «Номер», его
    # не трогаем; цена — числовой столбец правее текста
    numeric = [i for i in free() if _share(filled[i], number) >= 0.8]
    texts = [i for i in free() if i not in numeric]

    if texts and (price := next((i for i in numeric if i > min(texts)), None)) is not None:
        columns[price] = "price"

    long = max(texts, key=average, default=None)

    if long is not None and average(long) >= 80:
        columns[long] = "description"

    title = next((i for i in texts if i not in columns and average(i) >= 5), None)

    if title is None:
        return None

    columns[title] = "title"

    def worded(words: tuple[str, ...]) -> Any:  # noqa: ANN401
        return lambda v: any(w in v.lower() for w in words)

    for name, words in (("delivery_terms", _DELIVERY_WORDS), ("payment_terms", _PAYMENT_WORDS)):
        test = worded(words)
        scores = {i: _share(filled[i], test) for i in texts if i not in columns}
        best = max(scores, key=scores.__getitem__, default=None)

        if best is not None and _share(filled[best], test) >= 0.3:
            columns[best] = name
        elif (
            name == "payment_terms"
            and (delivery := next((i for i, n in columns.items() if n == "delivery_terms"), None))
            is not None
        ):
            # Оплату чаще пишут «По договорённости» — без примет; но она
            # обычно стоит сразу за поставкой, и хоть одна примета в столбце есть
            right = delivery + 1

            if right in scores and scores[right] > 0 and average(right) < 200:
                columns[right] = name

    sheet.header = first - 1
    sheet.sample = None

    for n, cells in rows:
        fields = {
            name: cells[i] for i, name in columns.items() if i < len(cells) and cells[i] != ""
        }
        offset = n - sheet.header
        sheet.rows[offset] = _Row(n, fields)

        if fields:
            sheet.last = offset

    found = ", ".join(f"{_letter(i)} — {_GUESSED[name]}" for i, name in sorted(columns.items()))

    if skipped_header is not None:
        return (
            f"Лист «{sheet.name}»: названия столбцов в строке {skipped_header} не узнаны — "
            f"столбцы определены по содержимому: {found}. Проверьте загруженное; чтобы не "
            "гадать, назовите столбцы, как в образце."
        )

    return (
        f"Лист «{sheet.name}»: нет строки с названиями столбцов — столбцы определены по "
        f"содержимому: {found}. Проверьте загруженное; чтобы не гадать, добавьте над "
        "таблицей строку с названиями, как в образце."
    )


def _looks_like_header(row: list[str], below: list[list[str]]) -> bool:
    """Строка названий: над столбцом чисел или длинных текстов — короткое слово."""
    for index, cell in enumerate(row):
        if cell == "" or len(cell) > 40:
            continue

        column = [cells[index] for cells in below if index < len(cells) and cells[index] != ""]

        if not column:
            continue

        numbers = _share(column, lambda v: re.fullmatch(r"[0-9][0-9\s.,]*", v.strip()) is not None)
        cell_is_number = re.fullmatch(r"[0-9][0-9\s.,]*", cell.strip()) is not None

        if numbers >= 0.8 and not cell_is_number:
            return True

        if sum(len(v) for v in column) / len(column) >= 80:
            return True

    return False


def _missing_header(sheet: _Sheet, expected: list[str]) -> str:
    """
    Почему на листе не нашлась шапка — по тому, что на нём лежит:
    пустой лист, вся шапка в одной ячейке через «;», или что записано
    в строке, похожей на шапку.
    """
    where = f"Лист «{sheet.name}»"

    if sheet.sample is None:
        return f"{where}: не найдено ни одной заполненной ячейки — лист пуст."

    number, cells = sheet.sample

    # CSV, пересохранённый в XLSX как текст: вся строка в первой ячейке
    if len(cells) == 1:
        for separator, label in (("\t", "табуляцию"), (";", "«;»"), (",", "«,»"), ("|", "«|»")):
            if separator in cells[0] and _columns(cells[0].split(separator), _EPOCH):
                return (
                    f"{where}, строка {number}: названия всех столбцов записаны в одной ячейке "
                    f"через {label}. Так бывает, когда CSV сохранён как XLSX без разбивки. "
                    f"Разнесите их по отдельным столбцам: в Excel «Данные → Текст по столбцам», "
                    f"разделитель {label}."
                )

    shown = ", ".join(f"«{_strimwidth(cell, 40)}»" for cell in cells[:6])
    more = f" и ещё {len(cells) - 6}" if len(cells) > 6 else ""

    return (
        f"{where}: не найдена строка с названиями столбцов. В строке {number} записано "
        f"{shown}{more} — ни одно не совпало с названием столбца. "
        f"Загрузка ждёт: {', '.join(expected)}."
    )


#: Начало отсчёта дат Excel по умолчанию (для разбора строк, где дат нет)
_EPOCH = datetime(1899, 12, 30)


# ── Загрузка одной книги ────────────────────────────────────────────


class _RowError(Exception):
    """RuntimeException у ListingWorkbookImport: строка не загружена, причина — в отчёт."""


@dataclass
class _Import:
    author_id: int
    replace: bool
    observer: access.Admin | None
    ip: str | None
    #: Компания для строк, где она не указана или не нашлась в справочнике:
    #: выбранная в окне загрузки, иначе служебная (service_company_id —
    #: ищется при первой надобности); нет ни той, ни другой — такие строки
    #: не создаются
    default_company: int | None = None
    service_looked_up: bool = False
    #: Тип новых объявлений, когда в книге нет столбца «Тип» или ячейка пуста
    default_type: str = "supply"
    #: «Сразу опубликовать»: новые строки — на витрину, без проверки
    publish: bool = False
    #: Ячейки, которые пришлось пропустить в текущей строке: незнакомая
    #: категория, город не из справочника — пропускаются, а объявление
    #: загружается. Сбрасывается перед каждой строкой
    skipped: list[str] = field(default_factory=list)
    #: Компания из строки, которую не нашли в справочнике
    unknown_company: str | None = None

    def run(self, content: bytes) -> WorkbookResult:
        result = _empty()
        sheets = self._sheets(content, result)
        roles = self._roles(sheets, result)

        if roles is None:
            return result

        master, translations = roles

        if not self._aligned(master, translations, result):
            return result

        # Строка называется с листом, только когда листов несколько
        named = bool(translations)

        with tempfile.TemporaryDirectory(prefix="listing-workbook-") as directory:
            photos = extract_photos(content, Path(directory), master.name)

            for offset, row in master.rows.items():
                if not row.fields and row.number not in photos:
                    continue

                result["rows"] += 1
                texts = {
                    locale: sheet.rows[offset].fields
                    for locale, sheet in translations.items()
                    if offset in sheet.rows
                }
                where = (
                    f"Лист «{master.name}», строка {row.number}"
                    if named
                    else f"Строка {row.number}"
                )
                stored: list[str] = []
                obsolete: list[str | None] = []
                # Пропуски — свои у каждой строки: иначе заметка второй
                # строки повторялась бы у всех следующих
                self.skipped = []

                try:
                    with transaction.atomic():
                        exists, saved = self._save_row(
                            row.fields, texts, photos.get(row.number, []), stored, obsolete
                        )
                except _RowError as error:
                    image_store.delete(*stored)
                    result["errors"].append(f"{where}: {error}")

                    continue
                except Exception:
                    log.exception("Строка книги товаров не загрузилась: %s", where)
                    image_store.delete(*stored)
                    result["errors"].append(f"{where}: строку не удалось загрузить")

                    continue

                # Прежние фотографии (replace) — с диска после записи строки
                image_store.delete(*obsolete)
                result["updated" if exists else "created"] += 1
                result["photos"] += saved

                # Что в строке пропущено — в отчёт: строка загружена,
                # но человек должен знать, чего в ней не хватает
                result["notes"].extend(f"{where}: {skip}" for skip in self.skipped)

        return result

    # ── Листы ────────────────────────────────────────────────────────

    def _sheets(self, content: bytes, result: WorkbookResult) -> list[_Sheet]:
        """
        Листы книги с разобранными строками: первый видимый и все с
        узнаваемым именем языка. Скрытые не читаются вовсе, прочие —
        справочники и заметки — пропускаются с заметкой в отчёте.
        """
        from openpyxl import load_workbook

        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as error:
            raise UnreadableWorkbookError("Файл не открывается как книга Excel (.xlsx).") from error

        sheets: list[_Sheet] = []

        try:
            epoch: datetime = book.epoch
            pages: list[Any] = list(book.worksheets)

            for page in pages:
                name = str(page.title)

                if page.sheet_state in ("hidden", "veryHidden"):
                    result["notes"].append(f"Скрытый лист «{name}» пропущен.")

                    continue

                locale = sheet_locale(name)

                if locale is None and sheets:
                    result["notes"].append(
                        f"Лист «{name}» пропущен: имя вкладки не узнано как язык. "
                        "Языковые вкладки называются «Русский», «English», «O‘zbekcha», «中文», "
                        "«Türkçe» — как в образце."
                    )

                    continue

                sheet = _Sheet(name, locale)
                _read_sheet(sheet, _values(page), epoch)
                sheets.append(sheet)
        except Exception as error:
            raise UnreadableWorkbookError("Книгу Excel не удалось прочитать.") from error
        finally:
            book.close()

        return sheets

    def _roles(
        self, sheets: list[_Sheet], result: WorkbookResult
    ) -> tuple[_Sheet, dict[str, _Sheet]] | None:
        """Какой лист главный и какие — переводы."""
        if not sheets:
            return None

        by_locale: dict[str, _Sheet] = {}

        for sheet in sheets:
            if sheet.locale is None:
                continue

            # Второй лист на том же языке не отменяет книгу: берётся первый
            if sheet.locale in by_locale:
                result["notes"].append(
                    f"Два листа на одном языке: «{by_locale[sheet.locale].name}» и "
                    f"«{sheet.name}». Взят первый, «{sheet.name}» пропущен."
                )

                continue

            by_locale[sheet.locale] = sheet

        master = by_locale.get("ru")

        # Русского по имени нет: главный — первый лист. Подписан другим
        # языком — всё равно он: цены и фотографии берутся с него, а книга
        # не отклоняется из-за имени вкладки
        if master is None:
            master = sheets[0]

            if master.locale is not None:
                result["notes"].append(
                    f"В книге нет листа «Русский» — главным взят первый лист «{master.name}»: "
                    "цены, категории и фотографии берутся с него."
                )
                by_locale.pop(master.locale, None)

        # Шапки нет — столбцы угадываются по содержимому (_guess_columns);
        # не вышло и тогда — объяснение, что лежит на листе
        if master.header is None and (note := _guess_columns(master)) is not None:
            result["notes"].append(note)

        if master.header is None:
            result["errors"].append(
                _missing_header(master, [h for h in HEADERS if h != "Фото"])
                + " Скачайте образец книги — в нём столбцы названы так, как их ждёт загрузка."
            )

            return None

        by_locale.pop("ru", None)

        return master, by_locale

    def _aligned(
        self, master: _Sheet, translations: dict[str, _Sheet], result: WorkbookResult
    ) -> bool:
        """
        Строки листов совпадают — иначе переводы съедут: и число строк,
        и напротив каждой заполненной строки перевода — заполненная русская.
        Лист перевода, который не сходится с русским или без шапки, не
        отменяет книгу: пропускается с заметкой, товары загружаются без
        этого языка.
        """
        for locale, sheet in list(translations.items()):
            if sheet.header is None:
                result["notes"].append(
                    _missing_header(
                        sheet, ["Заголовок", "Описание", "Условия поставки", "Условия оплаты"]
                    )
                    + " Лист перевода пропущен."
                )
                del translations[locale]

                continue

            if sheet.last == 0:
                result["notes"].append(f"Лист «{sheet.name}» пуст — переводов на этот язык нет.")
                del translations[locale]

                continue

            if sheet.last != master.last:
                result["notes"].append(
                    f"Лист «{sheet.name}»: строк с данными {sheet.last}, на русском листе "
                    f"{master.last}. Переводы связаны по порядку строк, при разном числе строк "
                    "они разъехались бы по чужим товарам, поэтому лист пропущен — товары "
                    "загружены без этого языка. Если перевода нет, оставьте строку пустой, "
                    "но не удаляйте её."
                )
                del translations[locale]

                continue

            for offset, row in sheet.rows.items():
                opposite = master.rows.get(offset)

                if row.fields and (opposite is None or not opposite.fields):
                    result["notes"].append(
                        f"Лист «{sheet.name}», строка {row.number}: заполнена, а на русском "
                        f"листе строка {opposite.number if opposite else row.number} пуста — "
                        "строки разъехались. Лист пропущен, товары загружены без этого языка."
                    )
                    del translations[locale]

                    break

        return True

    # ── Строка ───────────────────────────────────────────────────────

    def _save_row(
        self,
        fields: dict[str, str],
        texts: dict[str, dict[str, str]],
        files: list[Path],
        stored: list[str],
        obsolete: list[str | None],
    ) -> tuple[bool, int]:
        """Сохранить объявление строки; вернуть (было ли оно, сколько фото добавлено)."""
        original = self._resolve(fields)
        exists = "id" in original
        listing = dict(original)

        self._fill(listing, exists, fields)
        self._translate(listing, texts)
        listing = self._save(original, listing) if exists else self._insert(listing)

        if not _trim(str(listing.get("slug") or "")):
            self._quietly(listing, {"slug": _make_slug(str(listing["title"]), listing["id"])})

        return exists, self._attach(int(listing["id"]), files, stored, obsolete)

    def _resolve(self, fields: dict[str, str]) -> dict[str, Any]:
        """
        Объявление строки: по «Номеру»; номера нет в базе или он пуст —
        по заголовку и компании; иначе новое (только с компанией). Новое —
        без id.
        """
        # (int) у PHP не выходит за PHP_INT_MAX
        number = min(int(re.sub(r"[^0-9]", "", fields.get("id", "")) or 0), 2**63 - 1)

        if number > 0:
            found = _rows("select * from listings where id = %s and deleted_at is null", [number])

            if found:
                return found[0]

            # Номер устарел (объявление удалено) или опечатка — строка не
            # пропадает: ищется по заголовку и компании, как без номера
            self.skipped.append(
                f"объявление № {number} не найдено — строка загружена по заголовку и компании"
            )

        title = _trim(fields.get("title", ""))

        if title == "":
            raise _RowError("не заполнен заголовок")

        # Длинный заголовок сохранён обрезанным (_fits) — и искать его
        # надо таким же, иначе повторная загрузка заводит дубль
        title = _cut(title, MAX_LENGTH.get("title") or len(title))

        company = self._company_id(fields, fallback=True)
        query = "select * from listings where title = %s and deleted_at is null"
        params: list[Any] = [title]

        if company is not None:
            query += " and company_id = %s"
            params.append(company)

        found = _rows(query + " limit 2", params)

        # Без компании один заголовок может стоять у разных продавцов:
        # править первый попавшийся — значит менять цену чужого товара
        if len(found) > 1:
            raise _RowError(
                f"заголовок «{title}» есть у нескольких объявлений — укажите «Номер» или «Компанию»"
            )

        if found:
            return found[0]

        # Заявку раньше загрузили без компании (она у служебной), а теперь
        # в строке настоящая компания — заявка переходит к ней, а не
        # заводится второй раз
        fallback = self._fallback() if company is not None else None

        if company is not None and fallback is not None and company != fallback:
            parked = _rows(
                "select * from listings where title = %s and company_id = %s "
                "and deleted_at is null limit 2",
                [title, fallback],
            )

            if len(parked) == 1:
                self.skipped.append(
                    f"объявление передано от «{company_name(fallback)}» компании "
                    f"«{company_name(company)}»"
                )

                return parked[0]

        if company is None:
            raise _RowError(
                f"компания «{self.unknown_company}» не найдена в справочнике, а без компании "
                "новое объявление не создать. Добавьте её в «Компании» или выберите "
                "«Компанию для строк без компании» в окне загрузки"
                if self.unknown_company is not None
                else "не указана компания, а без неё новое объявление не создать. Заполните "
                "столбец «Компания» или выберите «Компанию для строк без компании» в окне загрузки"
            )

        return {"company_id": company}

    def _fallback(self) -> int | None:
        """Компания из окна загрузки, иначе служебная — одна на всю книгу."""
        if self.default_company is None and not self.service_looked_up:
            self.service_looked_up = True
            self.default_company = service_company_id()

        return self.default_company

    def _company_id(self, fields: dict[str, str], *, fallback: bool = False) -> int | None:
        """
        Компания строки из справочника. fallback — для поиска и нового
        объявления: пустая или ненайденная компания заменяется компанией
        из окна загрузки. У существующего объявления продавец без явной
        компании в строке не меняется.
        """
        name = _trim(fields.get("company", ""))
        self.unknown_company = None

        if name == "":
            return self._fallback() if fallback else None

        found = company_id(name)

        if found is None:
            if fallback and self._fallback() is not None:
                if f"компания «{name}»" not in " ".join(self.skipped):
                    self.skipped.append(
                        f"компания «{name}» не найдена в справочнике — объявление отнесено "
                        f"к «{company_name(self._fallback())}»"
                    )

                return self._fallback()

            # У нового объявления без компании нет продавца — строку это
            # остановит (_resolve); у существующего продавец уже есть
            self.unknown_company = name

            if f"компания «{name}»" not in " ".join(self.skipped):
                self.skipped.append(
                    f"компания «{name}» не найдена в справочнике — продавец остался прежним"
                )

        return found

    def _fill(self, listing: dict[str, Any], exists: bool, fields: dict[str, str]) -> None:
        """ListingWorkbookImport::fill: поля строки — в объявление; пустая ячейка не стирает."""
        if not exists:
            listing["user_id"] = self.author_id
            # Без ячейки «Тип» — по заголовку («Куплю…», «Продам…»), и только
            # если он ничего не говорит — тип из окна загрузки. Раньше всё шло
            # из окна, где по умолчанию «Предложение», и заявки «Куплю…» из
            # книг с других площадок становились предложениями
            listing["type"] = type_from_title(_trim(fields.get("title", ""))) or self.default_type
            # Ждёт проверки: публикует администратор из списка — или сразу,
            # галочкой «Сразу опубликовать» в окне загрузки
            listing["status"] = "active" if self.publish else "moderation"
        elif self.publish and listing.get("status") == "moderation":
            # Ждавшее проверки с прошлой загрузки публикуется той же галочкой;
            # отклонённое и снятое — нет: это решение модератора
            listing["status"] = "active"

        # Объявление, которое ведут книгой, живёт по правилам книги
        listing["source"] = "import"

        if (company := self._company_id(fields)) is not None:
            listing["company_id"] = company

        for plain in ("title", "description", "unit", "delivery_terms", "payment_terms"):
            if _trim(fields.get(plain, "")) != "":
                listing[plain] = self._fits(plain, _trim(fields[plain]))

        if _trim(fields.get("category_id", "")) != "":
            category, note = listing_category(fields["category_id"])

            if category is None:
                self.skipped.append(
                    f"категория «{_trim(fields['category_id'])}» не найдена в каталоге — "
                    "ячейка пропущена"
                )
            else:
                listing["category_id"] = category

                if note is not None:
                    self.skipped.append(note)

        if _trim(fields.get("city_id", "")) != "":
            city = city_id(fields["city_id"])

            if city is None:
                self.skipped.append(
                    f"город «{_trim(fields['city_id'])}» не найден в справочнике — ячейка пропущена"
                )
            else:
                listing["city_id"] = city

        if _trim(fields.get("type", "")) != "":
            kind = type_from_cell(fields["type"])

            if kind is None:
                # Непонятное слово — не «Предложение» молча, как раньше:
                # тип остаётся прежним (у новой строки — по заголовку или
                # из окна), а в отчёте — заметка
                self.skipped.append(
                    f"тип «{_trim(fields['type'])}» не распознан — оставлен "
                    f"«{'Запрос' if listing.get('type') == 'demand' else 'Предложение'}»"
                )
            else:
                listing["type"] = kind

        if _trim(fields.get("price", "")) != "":
            listing["price"] = importer.amount(fields["price"])
            listing["price_negotiable"] = listing["price"] is None

        if _trim(fields.get("currency", "")) != "":
            code = importer.currency(fields["currency"])

            if code not in CURRENCIES:
                self.skipped.append(
                    f"валюта «{_trim(fields['currency'])}» не поддерживается (можно: "
                    f"{', '.join(CURRENCIES)}) — оставлена прежняя"
                )
            else:
                listing["currency"] = code

        if _trim(fields.get("min_order", "")) != "":
            listing["min_order"] = int(importer.amount(fields["min_order"]) or 0) or None

        listing["currency"] = listing.get("currency") or "UZS"

        # Опубликованному нужен срок, иначе оно не попадает ни в один
        # список витрины; новому срок поставит публикация из админки
        if listing.get("status") == "active":
            now = _now()

            if listing.get("published_at") is None:
                listing["published_at"] = now

            if listing.get("expires_at") is None:
                listing["expires_at"] = now + timedelta(days=LIFETIME_DAYS)

    def _translate(self, listing: dict[str, Any], texts: dict[str, dict[str, str]]) -> None:
        """
        Тексты с языковых листов — в переводные колонки. Пустая ячейка
        не стирает перевод: переводчик мог ещё не дойти до строки.
        """
        for locale, values in texts.items():
            for name in TRANSLATABLE:
                value = _trim(values.get(name, ""))

                if value == "":
                    continue

                column = f"{name}_i18n"
                translations = _translations(listing.get(column))
                translations[locale] = self._fits(name, value, locale)
                listing[column] = translations

    def _fits(self, name: str, value: str, locale: str | None = None) -> str:
        """
        Текст не длиннее, чем принимает форма: длинный обрезается по
        границе слова (хвост, а не смысл), обрезка — в отчёт.
        """
        limit = MAX_LENGTH.get(name)

        if limit is None or len(value) <= limit:
            return value

        self.skipped.append(
            _LABELS.get(name, name)
            + (f" на языке «{locale}»" if locale is not None else "")
            + f" длиннее {limit} символов — обрезано"
        )
        return _cut(value, limit)

    # ── Запись ───────────────────────────────────────────────────────

    def _insert(self, listing: dict[str, Any]) -> dict[str, Any]:
        """$listing->save() нового: событие saving (search_text), журнал created."""
        listing["search_text"] = _search_text(
            {
                "title": listing.get("title"),
                "description": listing.get("description"),
                "title_i18n": listing.get("title_i18n"),
            }
        )
        now = _stamp(_now())
        listing.update(updated_at=now, created_at=now)
        columns = list(listing)

        with allowed_writes("listings"), connection.cursor() as cursor:
            cursor.execute(
                f"insert into listings ({', '.join(columns)}) "
                f"values ({', '.join(['%s'] * len(columns))}) returning id",
                [_written(CASTS.get(c), v) for c, v in listing.items()],
            )
            row = cursor.fetchone()
            assert row is not None
            listing["id"] = int(row[0])

        self._journal("created", listing, {"after": _attributes(listing)})

        return listing

    def _save(self, original: dict[str, Any], listing: dict[str, Any]) -> dict[str, Any]:
        """$listing->save() найденного: изменившиеся поля, search_text, журнал updated."""
        listing["search_text"] = _search_text(listing)
        dirty = {
            column: listing[column]
            for column in listing
            if column != "updated_at"
            and not _cast_same(CASTS.get(column), original.get(column), listing[column])
        }

        if not dirty:
            return listing

        now = _now()
        self._update(listing["id"], {**dirty, "updated_at": now})
        before = {c: _written(CASTS.get(c), original.get(c)) for c in [*dirty, "updated_at"]}
        listing["updated_at"] = now
        self._journal(
            "updated",
            listing,
            {"before": before, "after": _attributes({**dirty, "updated_at": now})},
        )

        return listing

    def _quietly(self, listing: dict[str, Any], changes: dict[str, Any]) -> None:
        """saveQuietly(): без событий и журнала, метка времени — как у save()."""
        now = _stamp(_now())

        if _stamp(listing.get("updated_at")) != now:
            changes = {**changes, "updated_at": now}

        self._update(listing["id"], changes)
        listing.update(changes)

    @staticmethod
    def _update(listing_id: int, changes: dict[str, Any]) -> None:
        sets = ", ".join(f"{column} = %s" for column in changes)

        with allowed_writes("listings"), connection.cursor() as cursor:
            cursor.execute(
                f"update listings set {sets} where id = %s",
                [*(_written(CASTS.get(c), v) for c, v in changes.items()), listing_id],
            )

    def _journal(self, action: str, listing: dict[str, Any], changes: dict[str, Any]) -> None:
        """AuditObserver: только правки администратора."""
        if self.observer is None:
            return

        audit.record(
            connection,
            action=action,
            section="listings",
            actor=self.observer,
            subject_type="App\\Models\\Listing",
            subject_id=int(listing["id"]),
            subject_label=audit.label(listing, "Listing", listing["id"]),
            changes=changes,
            ip=self.ip,
        )

    def _attach(
        self, listing_id: int, files: list[Path], stored: list[str], obsolete: list[str | None]
    ) -> int:
        """
        Фотографии строки — объявлению: только если своих нет (или
        replace); не больше MAX_IMAGES; нечитаемая картинка пропускается.
        """
        if not files:
            return 0

        images = _rows(
            "select id, path, thumb_path from listing_images where listing_id = %s "
            "order by sort, id",
            [listing_id],
        )
        already = len(images)

        if already > 0:
            if not self.replace:
                return 0

            with allowed_writes("listing_images"), connection.cursor() as cursor:
                for image in images:
                    obsolete.extend((image["path"], image["thumb_path"]))
                    cursor.execute("delete from listing_images where id = %s", [image["id"]])

            already = 0

        saved = 0

        for file in files[:MAX_IMAGES]:
            try:
                paths = image_store.store_with_thumb(file.read_bytes(), f"listings/{listing_id}")
            except (image_store.UnreadableImageError, OSError):
                # Одна нечитаемая картинка не должна ронять строку
                continue

            stored.extend((paths["path"], paths["thumb_path"]))
            now = _stamp(_now())

            with allowed_writes("listing_images"), connection.cursor() as cursor:
                cursor.execute(
                    "insert into listing_images (path, thumb_path, sort, listing_id, updated_at, "
                    "created_at) values (%s, %s, %s, %s, %s, %s)",
                    [paths["path"], paths["thumb_path"], already + saved, listing_id, now, now],
                )

            saved += 1

        return saved


def _values(page: Any) -> Iterable[list[Any]]:  # noqa: ANN401
    """
    Строки листа, как их читает openspout с SHOULD_PRESERVE_EMPTY_ROWS:
    пустые строки на месте (иначе фотография из седьмой строки
    достанется пятому товару), размер листа из заголовка не в счёт,
    ячейка с ошибкой формулы (#N/A) — пустая.
    """
    page.reset_dimensions()

    for row in page.iter_rows():
        yield [None if cell.data_type == "e" else cell.value for cell in row]


def _translations(value: Any) -> dict[str, Any]:  # noqa: ANN401
    """$listing->{$column} ?? []: переводы из базы или уже набранные."""
    if isinstance(value, str):
        value = json.loads(value)

    if isinstance(value, list):
        return {str(index): text for index, text in enumerate(value)}

    return dict(value) if isinstance(value, dict) else {}


def _attributes(values: dict[str, Any]) -> dict[str, Any]:
    """Атрибуты модели для журнала: массивы — текстом JSON, цена — дробным числом."""
    return {
        column: float(value) if isinstance(value, Decimal) else _written(CASTS.get(column), value)
        for column, value in values.items()
    }


def _make_slug(title: str, key: int) -> str:
    """Listing::makeSlug: хвост из заголовка (не длиннее 60) и номер."""
    return slugify(title)[:60].rstrip() + f"-{key}"
