"""
Потенциальные клиенты — база отдела продаж для рассылок (раздел CRM
«Потенциальные клиенты», savdex/crm/prospects_admin.py).

Это не компании площадки, а те, кого на неё зовут: компании, физлица,
фрилансеры. Данные бывают неполными — у физлица нет ИНН, у кого-то
только телефон, — и запись всё равно заводится; пустое в списке видно
как «—». Нужно хотя бы одно из названия, контактного лица, телефона,
почты или ИНН (IDENTITY).

Загрузка (import_rows) — таблица Excel или CSV, заголовки на русском,
узбекском или английском (COLUMNS). Повторы склеиваются: та же запись —
по ИНН, иначе по телефону (последние 9 цифр), иначе по почте; у записи
без всего этого — по названию и контакту. Найденная обновляется
непустыми ячейками (пустая ничего не стирает, «откуда» остаётся
первым, заметка дописывается к прежней), счётчик рассылок не
сбрасывается. Совпадение с компанией,
уже заведённой на площадке, — отметка «на SavdEx».

Рассылка — двух видов:

- письмо с площадки (queue_email): получатели с почтой, не отписавшиеся
  и (по умолчанию) не перенесённые в лиды ставятся в очередь, письма
  уходят фоном — manage.py notify, по PER_PASS за проход (run); счётчик
  растёт у тех, кому письмо ушло. В письме — ссылка «отписаться»
  (и заголовок List-Unsubscribe): отписавшемуся письма больше не идут;
- касание вручную (mark): продавец написал в Telegram, позвонил — и
  отмечает это у выделенных, счётчик +1 сразу.

«В лиды» (to_lead) — лид с источником «Холодный контакт», контактами и
сводкой из базы в заметке; запись остаётся с отметкой и ссылкой на лид.
"""

from __future__ import annotations

import hashlib
import hmac
import html
import logging
import os
import re
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from typing import Any

from django.conf import settings
from django.db import connection, transaction
from django.db.models import F, QuerySet
from django.utils import timezone

from savdex import access, audit
from savdex.catalog import now
from savdex.crm.models import (
    MAILING_CHANNELS,
    PROSPECT_KINDS,
    Lead,
    Prospect,
    ProspectMailing,
    ProspectRecipient,
)
from savdex.data.company_import import humanize
from savdex.tenders import importer
from savdex.web import mail

log = logging.getLogger("savdex.prospects")

SECTION = "prospects"
MODEL = "App\\Models\\Crm\\Prospect"
MAILING_MODEL = "App\\Models\\Crm\\ProspectMailing"

#: Поля записи, которые заливаются и правятся
FIELDS = (
    "kind",
    "name",
    "tin",
    "contact_person",
    "phone",
    "email",
    "city",
    "industry",
    "website",
    "source",
    "note",
)

#: Без хотя бы одного из них записи нет: не за что зацепиться
IDENTITY = ("name", "contact_person", "phone", "email", "tin")

#: Писем за проход фонового обработчика (раз в минуту): почтовые сервисы
#: режут тех, кто шлёт тысячи разом, и письма уходят в спам
PER_PASS = int(os.environ.get("SAVDEX_PROSPECT_MAILS_PER_MINUTE") or 30)

#: «Отправляется» дольше этого — проход упал посреди письма. Повторять
#: нельзя (письмо могло уйти), строка — «не ушло»
STALE = timedelta(minutes=30)

#: Метки в теме и тексте письма
PLACEHOLDERS = {
    "{имя}": "контактное лицо, иначе название, иначе «коллеги»",
    "{компания}": "название (пусто, если его нет)",
}


# ── Ключи для склейки повторов ─────────────────────────────────────


def _digits(value: str | None) -> str:
    return re.sub(r"\D", "", value or "")


def tin_key(tin: str | None) -> str | None:
    digits = _digits(tin)

    return "t:" + digits if len(digits) >= 5 else None


def phone_key(phone: str | None) -> str | None:
    """Первый телефон ячейки, последние 9 цифр: +998 90…, 8 90…, 90… — один номер."""
    first = re.split(r"[,;/\n]", phone or "", maxsplit=1)[0]
    digits = _digits(first)

    return "p:" + digits[-9:] if len(digits) >= 7 else None


def email_key(email: str | None) -> str | None:
    value = (email or "").strip().lower()

    return "e:" + value if "@" in value else None


def keys(record: Mapping[str, Any]) -> list[str]:
    """
    Чем запись узнаётся: ИНН, телефон, почта. У записи без них — название
    с контактом (физлицо из одного имени тоже не должно двоиться).
    """
    found = [
        key
        for key in (
            tin_key(record.get("tin")),
            phone_key(record.get("phone")),
            email_key(record.get("email")),
        )
        if key
    ]

    if not found:
        name = importer.normalize(record.get("name") or "")
        contact = importer.normalize(record.get("contact_person") or "")

        if name or contact:
            found.append(f"n:{name}|{contact}")

    return found


def find(record: Mapping[str, Any], *, exclude: int | None = None) -> Prospect | None:
    """Уже заведённая запись с тем же ИНН, телефоном или почтой (для формы)."""
    from django.db.models import Q

    from savdex import search

    wanted = set(keys(record))

    if not wanted:
        return None

    tin = _digits(record.get("tin"))
    email = (record.get("email") or "").strip()
    phone = phone_key(record.get("phone"))
    name = (record.get("name") or record.get("contact_person") or "").strip()
    near = Q(pk__in=[])

    # Грубый отбор в базе, точное сравнение ключей — ниже
    if len(tin) >= 5:
        near |= Q(tin__contains=tin[-5:])

    if email:
        near |= Q(email__iexact=email)

    if phone:
        # Последние четыре цифры, между ними — что угодно: «43-21», «4321»
        near |= Q(phone__regex=r"\D*".join(phone[-4:]))

    if name and not (len(tin) >= 5 or email or phone):
        near |= Q(tin__isnull=True, phone__isnull=True, email__isnull=True) & (
            search.contains("name", name) | search.contains("contact_person", name)
        )

    candidates = Prospect.objects.filter(near)

    if exclude is not None:
        candidates = candidates.exclude(pk=exclude)

    found: Prospect | None = next(
        (p for p in candidates.order_by("id")[:500] if wanted & set(keys(_record(p)))), None
    )

    return found


def _record(prospect: Prospect) -> dict[str, Any]:
    return {name: getattr(prospect, name) for name in FIELDS}


# ── Компании площадки ──────────────────────────────────────────────


class Companies:
    """Компании SavdEx по ИНН, почте и телефону — для отметки «на SavdEx»."""

    def __init__(self) -> None:
        self.by_key: dict[str, int] = {}

        with connection.cursor() as cursor:
            cursor.execute(
                "select id, tin, email, phone from companies where deleted_at is null order by id"
            )

            for pk, tin, email, phone in cursor.fetchall():
                for key in (tin_key(tin), email_key(email), phone_key(phone)):
                    if key:
                        self.by_key.setdefault(key, int(pk))

    def find(self, record: Mapping[str, Any]) -> int | None:
        for key in (
            tin_key(record.get("tin")),
            email_key(record.get("email")),
            phone_key(record.get("phone")),
        ):
            if key and key in self.by_key:
                return self.by_key[key]

        return None


def savdex_company(record: Mapping[str, Any]) -> int | None:
    """Одна запись (форма): та же компания на площадке, если есть."""
    tin = _digits(record.get("tin"))
    email = (record.get("email") or "").strip().lower()
    phone = phone_key(record.get("phone"))
    checks = []

    if len(tin) >= 5:
        checks.append(("regexp_replace(coalesce(tin, ''), '\\D', '', 'g') = %s", tin))

    if email:
        checks.append(("lower(email) = %s", email))

    if phone:
        checks.append(
            ("right(regexp_replace(coalesce(phone, ''), '\\D', '', 'g'), 9) = %s", phone[2:])
        )

    with connection.cursor() as cursor:
        for where, value in checks:
            cursor.execute(
                f"select id from companies where deleted_at is null and {where} "
                "order by id limit 1",
                [value],
            )
            row = cursor.fetchone()

            if row:
                return int(row[0])

    return None


# ── Загрузка ────────────────────────────────────────────────────────


@dataclass(frozen=True)
class Column:
    name: str
    #: Заголовок в образце
    label: str
    hint: str
    #: Заголовки, которые узнаются (нормализованные: строчные, «е» вместо «ё»)
    aliases: tuple[str, ...]


COLUMNS = (
    Column(
        "name",
        "Название / ФИО",
        "компания или ФИО физлица; капс — «Каждое Слово С Заглавной»",
        ("название / фио", "название/фио", "название", "наименование", "компания",
         "название компании", "организация", "фирма", "клиент", "контрагент",
         "name", "company", "company name", "client", "nomi", "kompaniya", "tashkilot"),
    ),
    Column(
        "kind",
        "Кто это",
        "компания, физлицо или фрилансер; пусто — не указано",
        ("кто это", "кто", "тип", "тип клиента", "вид", "форма", "лицо",
         "kind", "type", "client type", "turi", "shaxs"),
    ),
    Column(
        "tin",
        "ИНН",
        "у физлица и фрилансера — можно пусто",
        ("инн", "налоговый номер", "стир", "пинфл", "tin", "inn", "stir", "tax id", "jshshir"),
    ),
    Column(
        "contact_person",
        "Контактное лицо",
        "кому писать: «Здравствуйте, {имя}!»",
        ("контактное лицо", "контакт", "фио", "имя", "ответственный", "представитель",
         "директор", "руководитель", "contact", "contact person", "contact name", "person",
         "kontakt", "masul shaxs", "mas'ul shaxs", "ism", "fio", "f.i.o."),
    ),
    Column(
        "phone",
        "Телефон",
        "можно несколько через запятую",
        ("телефон", "тел", "тел.", "телефоны", "номер", "номер телефона", "мобильный",
         "контактный телефон", "phone", "telephone", "mobile", "tel", "telefon", "raqam",
         "telefon raqami"),
    ),
    Column(
        "email",
        "Почта",
        "на неё уходят письма; адрес с ошибкой не записывается",
        ("почта", "эл. почта", "эл.почта", "электронная почта", "email", "e-mail", "mail",
         "емейл", "имейл", "pochta", "elektron pochta"),
    ),
    Column(
        "city",
        "Город",
        "",
        ("город", "регион", "область", "населенный пункт", "адрес", "city", "region",
         "address", "shahar", "viloyat", "manzil"),
    ),
    Column(
        "industry",
        "Отрасль",
        "чем занимается",
        ("отрасль", "сфера", "сфера деятельности", "направление", "категория",
         "вид деятельности", "деятельность", "чем занимается", "industry", "category",
         "activity", "soha", "faoliyat", "yonalish"),
    ),
    Column(
        "website",
        "Сайт",
        "",
        ("сайт", "веб-сайт", "вебсайт", "website", "site", "web", "url", "sayt"),
    ),
    Column(
        "note",
        "Заметка",
        "",
        ("заметка", "комментарий", "примечание", "note", "comment", "notes", "izoh"),
    ),
    Column(
        "source",
        "Откуда",
        "где нашли; пусто — то, что указано в окне загрузки",
        ("откуда", "источник", "source", "manba"),
    ),
)  # fmt: skip

#: «Кто это» в ячейке → код (сравнение без пробелов, точек и дефисов)
_KINDS = {
    "company": ("компания", "юрлицо", "юридическоелицо", "организация", "фирма", "ооо",
                "company", "legal", "yuridikshaxs", "kompaniya", "tashkilot"),
    "person": ("физлицо", "физ", "физическоелицо", "частноелицо", "частник", "person",
               "individual", "jismoniyshaxs", "jismoniy"),
    "freelancer": ("фрилансер", "фриланс", "самозанятый", "freelancer", "freelance",
                   "frilanser"),
}  # fmt: skip


def kind_of(value: Any) -> str | None:  # noqa: ANN401
    """«Физ. лицо» → person; незнакомое — None."""
    bare = re.sub(r"[\s.\-]", "", importer.normalize(value))

    if not bare:
        return None

    if bare in PROSPECT_KINDS:
        return bare

    return next((code for code, words in _KINDS.items() if bare in words), None)


def column_map(headers: Iterable[str]) -> dict[str, str]:
    """Поле → заголовок файла; каждый заголовок — одному полю."""
    headers = [h for h in headers if h]
    mapping: dict[str, str] = {}
    taken: set[str] = set()

    for column in COLUMNS:
        aliases = {importer.normalize(a) for a in (column.label, *column.aliases)}

        for header in headers:
            if header not in taken and importer.matches(header, aliases):
                mapping[column.name] = header
                taken.add(header)
                break

    return mapping


def _cast(raw: Mapping[str, Any], mapping: Mapping[str, str]) -> tuple[dict[str, Any], list[str]]:
    """Строка таблицы → поля записи и заметки о том, что не записано."""

    def cell(name: str) -> Any:  # noqa: ANN401
        return raw.get(mapping[name]) if name in mapping else None

    data: dict[str, Any] = {}
    notes: list[str] = []

    if name := importer.text(cell("name"), 190):
        data["name"] = humanize(name)

    if (kind_cell := importer.text(cell("kind"))) is not None:
        if kind := kind_of(kind_cell):
            data["kind"] = kind
        else:
            notes.append(f"«{kind_cell[:40]}» в «Кто это» не узнано — пусто")

    if tin := importer.first(cell("tin"), 30):
        data["tin"] = re.sub(r"\s", "", tin)[:20]

    if email_cell := importer.text(cell("email")):
        if email := importer.email(email_cell):
            data["email"] = email
        else:
            notes.append(f"почта «{email_cell[:60]}» не похожа на адрес — не записана")

    limits = {
        "contact_person": 160,
        "phone": 100,
        "city": 120,
        "industry": 190,
        "source": 120,
        "note": 5000,
    }

    for field_name, limit in limits.items():
        if value := importer.text(cell(field_name), limit):
            data[field_name] = value

    if website := importer.first(cell("website"), 255):
        data["website"] = website

    return data, notes


@dataclass
class Report:
    """Итог загрузки; номер строки — как в Excel (заголовок — первая)."""

    created: int = 0
    updated: int = 0
    unchanged: int = 0
    #: Строки, слитые с другой строкой того же файла
    merged: int = 0
    #: Записи из файла, которые уже есть на площадке
    on_savdex: int = 0
    failed: list[tuple[int, str]] = field(default_factory=list)
    notes: list[tuple[int, str]] = field(default_factory=list)
    unknown_headers: list[str] = field(default_factory=list)
    #: В файле нет ни одного столбца, по которому узнаётся запись
    missing: bool = False

    def summary(self) -> str:
        if self.missing:
            return (
                "В файле не нашлось столбцов с названием, контактом, телефоном, почтой "
                "или ИНН — ничего не загружено."
            )

        text = (
            f"Загружено: новых {self.created}, обновлено {self.updated}, "
            f"без изменений {self.unchanged}"
        )

        if self.merged:
            text += f", повторов внутри файла {self.merged}"

        if self.failed:
            text += f", пропущено {len(self.failed)}"

        return text + "."


def import_rows(
    table: list[dict[str, Any]],
    *,
    staff: access.Admin,
    source: str | None = None,
    kind: str | None = None,
    ip: str | None = None,
    file_name: str = "",
) -> Report:
    """
    Строки таблицы (importer.read_table) — в базу: новые заводятся,
    найденные обновляются. Всё одной транзакцией; строка журнала
    «Загрузка» с итогами.
    """
    report = Report()
    headers = list(table[0].keys()) if table else []
    mapping = column_map(headers)
    report.unknown_headers = [h for h in headers if h and h not in mapping.values()]

    if not any(name in mapping for name in IDENTITY):
        report.missing = True

        return report

    existing: dict[int, dict[str, Any]] = {}

    for row in Prospect.objects.values("id", "company_id", *FIELDS):
        existing[int(row["id"])] = row

    index: dict[str, dict[str, Any]] = {}

    for record in existing.values():
        for key in keys(record):
            index.setdefault(key, record)

    fresh: list[dict[str, Any]] = []
    touched: dict[int, dict[str, Any]] = {}
    seen: set[int] = set()

    for number, raw in enumerate(table, start=2):
        data, notes = _cast(raw, mapping)
        report.notes += [(number, note) for note in notes]

        if not any(data.get(name) for name in IDENTITY):
            if any(data.values()):
                report.failed.append(
                    (number, "нет ни названия, ни контактного лица, ни телефона, ни почты, ни ИНН")
                )

            continue

        if kind and not data.get("kind"):
            data["kind"] = kind

        if source and not data.get("source"):
            data["source"] = source

        target = next((index[key] for key in keys(data) if key in index), None)

        if target is None:
            record = {**dict.fromkeys(FIELDS), **data, "id": None, "company_id": None}
            fresh.append(record)

            for key in keys(record):
                index.setdefault(key, record)

            continue

        # Заметка дописывается к прежней, а не затирает её
        if data.get("note") and target.get("note"):
            data["note"] = (
                target["note"]
                if data["note"] in target["note"]
                else f"{target['note']}\n{data['note']}"
            )

        changes = {
            name: value
            for name, value in data.items()
            if value != target.get(name) and not (name == "source" and target.get("source"))
        }
        target.update(changes)

        for key in keys(target):
            index.setdefault(key, target)

        if target["id"] is None:
            report.merged += 1
        else:
            seen.add(int(target["id"]))

            if changes:
                touched.setdefault(int(target["id"]), {}).update(changes)

    companies = Companies()

    for record in fresh:
        record["company_id"] = companies.find(record)

    for pk, changes in touched.items():
        if existing[pk]["company_id"] is None and (company := companies.find(existing[pk])):
            changes["company_id"] = company

    stamp = now()

    with transaction.atomic():
        Prospect.objects.bulk_create(
            [
                Prospect(
                    **{name: record[name] for name in FIELDS},
                    company_id=record["company_id"],
                    created_by=staff.id,
                    created_at=stamp,
                    updated_at=stamp,
                )
                for record in fresh
            ],
            batch_size=500,
        )

        for pk, changes in touched.items():
            Prospect.objects.filter(pk=pk).update(**changes, updated_at=stamp)

    report.created = len(fresh)
    report.updated = len(touched)
    report.unchanged = len(seen - set(touched))
    report.on_savdex = sum(1 for r in fresh if r["company_id"]) + sum(
        1 for pk in touched if existing[pk].get("company_id") or touched[pk].get("company_id")
    )

    audit.record(
        connection,
        action="imported",
        section=SECTION,
        actor=staff,
        note=(
            f"{file_name or 'Файл'}: строк {len(table)}, новых {report.created}, "
            f"обновлено {report.updated}, пропущено {len(report.failed)}"
        ),
        ip=ip,
    )

    return report


# ── Образец файла ──────────────────────────────────────────────────

#: Строки образца: компания со всем, физлицо и фрилансер — с пробелами
SAMPLE_ROWS = (
    ("ООО «Мебель Плюс»", "Компания", "301234567", "Азиз Каримов", "+998 90 123-45-67",
     "info@mebelplus.uz", "Ташкент", "Мебель", "mebelplus.uz", "Просили позвонить в марте",
     "Выставка UzBuild 2026"),
    ("Иванов Сергей Петрович", "Физлицо", "", "", "+998 93 765-43-21", "",
     "Самарканд", "Строительство", "", "", ""),
    ("", "Фрилансер", "", "Дилноза", "", "dilnoza.design@gmail.com", "", "Дизайн интерьеров",
     "", "", "Instagram"),
)  # fmt: skip


def sample_workbook() -> bytes:
    """Образец .xlsx: заголовки с подсказками в примечаниях и три строки-примера."""
    import io

    from openpyxl import Workbook
    from openpyxl.cell.cell import Cell
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Потенциальные клиенты"
    sheet.append([column.label for column in COLUMNS])

    for row in SAMPLE_ROWS:
        sheet.append(list(row))

    widths = (30, 13, 13, 22, 22, 28, 14, 22, 18, 30, 24)

    for number, (column, width) in enumerate(zip(COLUMNS, widths, strict=True), start=1):
        head = sheet.cell(row=1, column=number)
        assert isinstance(head, Cell)
        head.font = Font(bold=True, color="FFFFFF")
        head.fill = PatternFill("solid", fgColor="1F4E79")
        head.alignment = Alignment(vertical="center", wrap_text=True)

        if column.hint:
            head.comment = Comment(column.hint, "SavdEx")

        sheet.column_dimensions[head.column_letter].width = width

    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    book.save(buffer)

    return buffer.getvalue()


# ── Рассылки ───────────────────────────────────────────────────────


def mail_ready() -> bool:
    """Почта площадки настроена — письма уйдут, а не лягут в файл журнала."""
    mailer = os.environ.get("MAIL_MAILER") or "log"

    return mail.configured() or (settings.DEBUG and mailer == "log")


@dataclass
class Audience:
    """Кому из выделенных уйдёт письмо и почему остальным — нет."""

    selected: int
    recipients: QuerySet[Prospect]
    count: int
    no_email: int
    unsubscribed: int
    converted: int

    def missed(self) -> str:
        """Кому из отмеченных письмо не уйдёт и почему — одной строкой."""
        parts = [
            f"{label} — {count}"
            for label, count in (
                ("без почты", self.no_email),
                ("«не писать»", self.unsubscribed),
                ("уже в лидах", self.converted),
            )
            if count
        ]

        return ", ".join(parts)


def email_audience(selected: QuerySet[Prospect], *, include_converted: bool = False) -> Audience:
    rows = selected.filter(deleted_at__isnull=True)
    with_email = rows.exclude(email__isnull=True).exclude(email="")
    allowed = with_email.filter(unsubscribed_at__isnull=True)
    recipients = allowed if include_converted else allowed.filter(converted_at__isnull=True)
    count = recipients.count()

    return Audience(
        selected=rows.count(),
        recipients=recipients,
        count=count,
        no_email=rows.count() - with_email.count(),
        unsubscribed=with_email.count() - allowed.count(),
        converted=0 if include_converted else allowed.count() - count,
    )


def queue_email(
    recipients: QuerySet[Prospect],
    *,
    subject: str,
    body: str,
    reply_to: str | None,
    staff: access.Admin,
    ip: str | None = None,
) -> ProspectMailing:
    """Письмо — в очередь: строка на каждого получателя, уходит фоном (run)."""
    stamp = now()

    with transaction.atomic():
        mailing = ProspectMailing(
            channel="email",
            subject=subject,
            body=body,
            reply_to=reply_to or None,
            sent_by_id=staff.id,
        )
        mailing.save()
        rows = list(recipients.order_by("id").values_list("id", "email"))
        ProspectRecipient.objects.bulk_create(
            [
                ProspectRecipient(
                    mailing=mailing,
                    prospect_id=pk,
                    email=email,
                    status="queued",
                    created_at=stamp,
                    updated_at=stamp,
                )
                for pk, email in rows
            ],
            batch_size=1000,
        )
        mailing.total = len(rows)

        if not rows:
            mailing.finished_at = stamp

        mailing.save()

    audit.record(
        connection,
        action="sent",
        section=SECTION,
        actor=staff,
        subject_type=MAILING_MODEL,
        subject_id=mailing.pk,
        subject_label=subject,
        note=f"Письмо в очереди: получателей {len(rows)}",
        ip=ip,
    )

    return mailing


def mark(
    selected: QuerySet[Prospect],
    *,
    channel: str,
    note: str | None,
    staff: access.Admin,
    ip: str | None = None,
) -> ProspectMailing:
    """Касание вручную (Telegram, звонок…): у выделенных счётчик +1 сразу."""
    stamp = now()

    with transaction.atomic():
        alive = selected.filter(deleted_at__isnull=True).order_by("id")
        ids = list(alive.values_list("id", flat=True))
        mailing = ProspectMailing(
            channel=channel,
            note=note or None,
            sent_by_id=staff.id,
            total=len(ids),
            sent=len(ids),
            finished_at=stamp,
        )
        mailing.save()
        ProspectRecipient.objects.bulk_create(
            [
                ProspectRecipient(
                    mailing=mailing,
                    prospect_id=pk,
                    status="sent",
                    sent_at=stamp,
                    created_at=stamp,
                    updated_at=stamp,
                )
                for pk in ids
            ],
            batch_size=1000,
        )
        Prospect.objects.filter(pk__in=ids).update(
            mailings_count=F("mailings_count") + 1, last_mailed_at=stamp
        )

    label = MAILING_CHANNELS.get(channel, channel)
    audit.record(
        connection,
        action="sent",
        section=SECTION,
        actor=staff,
        subject_type=MAILING_MODEL,
        subject_id=mailing.pk,
        subject_label=str(mailing),
        note=f"{label}, отмечено вручную: получателей {len(ids)}",
        ip=ip,
    )

    return mailing


def stop(mailing: ProspectMailing, *, staff: access.Admin, ip: str | None = None) -> int:
    """Остановить письмо: кому ещё не ушло — «пропущено»."""
    stamp = now()

    with transaction.atomic():
        stopped = ProspectRecipient.objects.filter(mailing=mailing, status="queued").update(
            status="skipped", error="рассылка остановлена", updated_at=stamp
        )
        _finish(stamp)

    if stopped:
        audit.record(
            connection,
            action="updated",
            section=SECTION,
            actor=staff,
            subject_type=MAILING_MODEL,
            subject_id=mailing.pk,
            subject_label=str(mailing),
            note=f"Рассылка остановлена: не отправлено {stopped}",
            ip=ip,
        )

    return stopped


# ── Письмо ─────────────────────────────────────────────────────────


def _app_url() -> str:
    return (os.environ.get("APP_URL") or "http://localhost").rstrip("/")


def unsubscribe_token(pk: int) -> str:
    return hmac.new(
        str(settings.SECRET_KEY).encode(),
        f"prospect-unsubscribe:{pk}".encode(),
        hashlib.sha256,
    ).hexdigest()[:32]


def unsubscribe_url(pk: int) -> str:
    return f"{_app_url()}/unsubscribe/{pk}/{unsubscribe_token(pk)}"


def personalize(text: str, prospect: Prospect) -> str:
    """{имя} и {компания} — из записи."""
    return text.replace("{имя}", prospect.contact_person or prospect.name or "коллеги").replace(
        "{компания}", prospect.name or ""
    )


_URL = re.compile(r"https?://[^\s<>\"']+")


def _paragraphs(text: str) -> str:
    """Текст письма → абзацы HTML: экранирование, ссылки, переносы строк."""
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text.replace("\r\n", "\n")) if b.strip()]
    out = []

    for block in blocks:
        escaped = html.escape(block)
        linked = _URL.sub(lambda m: f'<a href="{m.group(0)}">{m.group(0)}</a>', escaped)
        out.append(
            '<p style="margin:0 0 14px;font-size:15px;line-height:1.55;color:#1f2937;">'
            + linked.replace("\n", "<br>")
            + "</p>"
        )

    return "".join(out)


def render(mailing: ProspectMailing, prospect: Prospect) -> tuple[str, str, str]:
    """Тема, HTML и текст письма этому получателю."""
    subject = personalize(mailing.subject or "", prospect).strip() or "SavdEx"
    body = personalize(mailing.body or "", prospect).strip()
    link = unsubscribe_url(int(prospect.pk))
    text = (
        f"{body}\n\n—\nSavdEx — B2B-площадка Узбекистана, {_app_url()}\n"
        f"Больше не присылать письма / Xat yubormaslik: {link}\n"
    )
    page = (
        '<!doctype html><html><body style="margin:0;padding:24px;background:#f4f5f7;">'
        '<div style="max-width:600px;margin:0 auto;background:#ffffff;border-radius:8px;'
        "padding:28px 32px;font-family:-apple-system,BlinkMacSystemFont,Segoe UI,Roboto,"
        'Helvetica,Arial,sans-serif;">'
        + _paragraphs(body)
        + '</div><p style="max-width:600px;margin:16px auto 0;font-family:Arial,sans-serif;'
        'font-size:12px;line-height:1.5;color:#6b7280;text-align:center;">'
        f'SavdEx — B2B-площадка Узбекистана · <a href="{html.escape(_app_url())}" '
        'style="color:#6b7280;">savdex.uz</a><br>'
        f'<a href="{html.escape(link)}" style="color:#6b7280;">Больше не присылать письма'
        " / Xat yubormaslik</a></p></body></html>"
    )

    return subject, page, text


# ── Фоновая отправка (manage.py notify) ─────────────────────────────


@dataclass
class Pass:
    sent: int = 0
    failed: int = 0
    skipped: int = 0

    def __bool__(self) -> bool:
        return bool(self.sent or self.failed or self.skipped)

    def __str__(self) -> str:
        return f"отправлено {self.sent}, не ушло {self.failed}, пропущено {self.skipped}"


def _claim(stamp: datetime, limit: int) -> list[tuple[int, int, int]]:
    """Занять строки очереди (skip locked): два прохода одно письмо не отправят."""
    with transaction.atomic(), connection.cursor() as cursor:
        cursor.execute(
            "update crm_prospect_mailing_recipients set status = 'sending', updated_at = %s "
            "where id in (select id from crm_prospect_mailing_recipients where status = 'queued' "
            "order by id limit %s for update skip locked) "
            "returning id, mailing_id, prospect_id",
            [stamp.replace(tzinfo=None), limit],
        )
        rows = [(int(a), int(b), int(c)) for a, b, c in cursor.fetchall()]

    return sorted(rows)


def _finish(stamp: datetime) -> None:
    """Очередь письма пуста — рассылка закончена; итоги — по строкам получателей."""
    with connection.cursor() as cursor:
        cursor.execute(
            "update crm_prospect_mailings m set finished_at = %s, updated_at = %s, "
            "sent = (select count(*) from crm_prospect_mailing_recipients r "
            "        where r.mailing_id = m.id and r.status = 'sent'), "
            "failed = (select count(*) from crm_prospect_mailing_recipients r "
            "          where r.mailing_id = m.id and r.status = 'failed') "
            "where m.finished_at is null and not exists (select 1 from "
            "crm_prospect_mailing_recipients r where r.mailing_id = m.id "
            "and r.status in ('queued', 'sending'))",
            [stamp.replace(tzinfo=None), stamp.replace(tzinfo=None)],
        )


def run(stamp: datetime | None = None, limit: int = PER_PASS) -> Pass:
    """Проход: до limit писем из очереди, счётчики — тем, кому ушло."""
    stamp = stamp or now()
    report = Pass()

    # Прерванные посреди письма: повторить нельзя — могло и уйти
    with transaction.atomic():
        report.failed += ProspectRecipient.objects.filter(
            status="sending", updated_at__lt=stamp - STALE
        ).update(status="failed", error="отправка прервана", updated_at=stamp)

    mailings: dict[int, ProspectMailing] = {}

    for pk, mailing_id, prospect_id in _claim(stamp, limit):
        mailing = mailings.get(mailing_id) or ProspectMailing.objects.get(pk=mailing_id)
        mailings[mailing_id] = mailing
        prospect = Prospect.everything.filter(pk=prospect_id).first()
        reason = None

        if prospect is None or prospect.deleted_at is not None:
            reason = "запись удалена"
        elif prospect.unsubscribed_at is not None:
            reason = "отписался"
        elif not prospect.email:
            reason = "нет почты"

        if reason or prospect is None or not prospect.email:
            ProspectRecipient.objects.filter(pk=pk).update(
                status="skipped", error=reason, updated_at=now()
            )
            report.skipped += 1
            continue

        subject, body_html, body_text = render(mailing, prospect)
        link = unsubscribe_url(int(prospect.pk))
        ok = mail.send(
            prospect.email,
            subject,
            body_html,
            body_text,
            reply_to=mailing.reply_to,
            headers={
                "List-Unsubscribe": f"<{link}>",
                "List-Unsubscribe-Post": "List-Unsubscribe=One-Click",
            },
        )
        sent_at = now()

        if ok:
            with transaction.atomic():
                ProspectRecipient.objects.filter(pk=pk).update(
                    status="sent", email=prospect.email, sent_at=sent_at, updated_at=sent_at
                )
                Prospect.everything.filter(pk=prospect.pk).update(
                    mailings_count=F("mailings_count") + 1, last_mailed_at=sent_at
                )
                ProspectMailing.objects.filter(pk=mailing_id).update(sent=F("sent") + 1)
            report.sent += 1
        else:
            ProspectRecipient.objects.filter(pk=pk).update(
                status="failed",
                email=prospect.email,
                error="почтовик не принял письмо",
                updated_at=sent_at,
            )
            ProspectMailing.objects.filter(pk=mailing_id).update(failed=F("failed") + 1)
            report.failed += 1

    _finish(now())

    return report


# ── Отписка ────────────────────────────────────────────────────────


def unsubscribe(pk: int, token: str) -> Prospect | None:
    """Ссылка из письма: верный ключ — больше не писать. Неверный — None."""
    if not hmac.compare_digest(token, unsubscribe_token(pk)):
        return None

    prospect = Prospect.everything.filter(pk=pk).first()

    if prospect is not None and prospect.unsubscribed_at is None:
        stamp = now()
        Prospect.everything.filter(pk=pk).update(unsubscribed_at=stamp)
        ProspectRecipient.objects.filter(prospect_id=pk, status="queued").update(
            status="skipped", error="отписался", updated_at=stamp
        )
        prospect.unsubscribed_at = stamp
        audit.record(
            connection,
            action="updated",
            section=SECTION,
            actor=None,
            subject_type=MODEL,
            subject_id=pk,
            subject_label=prospect.label(),
            note="Отписался по ссылке из письма",
        )

    return prospect


# ── В лиды ─────────────────────────────────────────────────────────


def _lead_note(prospect: Prospect) -> str:
    lines = [f"Из базы потенциальных клиентов, №{prospect.pk}."]
    facts = (
        ("Кто", PROSPECT_KINDS.get(prospect.kind or "", "")),
        ("Название", prospect.name if prospect.contact_person else ""),
        ("ИНН", prospect.tin),
        ("Телефоны", prospect.phone if re.search(r"[,;/]", prospect.phone or "") else ""),
        ("Город", prospect.city),
        ("Отрасль", prospect.industry),
        ("Сайт", prospect.website),
        ("Откуда", prospect.source),
    )
    lines += [f"{label}: {value}" for label, value in facts if value]

    if prospect.mailings_count:
        last = (
            timezone.localtime(prospect.last_mailed_at).strftime("%d.%m.%Y")
            if prospect.last_mailed_at
            else "—"
        )
        lines.append(f"Рассылок: {prospect.mailings_count}, последняя {last}")

    if prospect.note:
        lines.append(f"Заметка: {prospect.note}")

    return "\n".join(lines)


def to_lead(prospect: Prospect, *, staff: access.Admin, ip: str | None = None) -> Lead | None:
    """
    Лид с источником «Холодный контакт»: контакты из записи, сводка — в
    заметке, ответственный — тот, кто переносит. Уже в лидах — None.
    """
    if prospect.converted_at is not None:
        return None

    email = prospect.email if prospect.email and len(prospect.email) <= 160 else None
    stamp = now()

    with transaction.atomic():
        lead = Lead(
            title=prospect.label()[:200],
            source="outbound",
            company_id=prospect.company_id,
            contact_name=(prospect.contact_person or prospect.name or None),
            contact_phone=importer.first(prospect.phone, 40),
            contact_email=email,
            owner_id=staff.id if staff.can("leads.edit") else None,
            status="new",
            note=_lead_note(prospect),
        )
        lead.save()
        Prospect.objects.filter(pk=prospect.pk).update(
            lead_id=lead.pk, converted_at=stamp, updated_at=stamp
        )
        prospect.lead_id = lead.pk
        prospect.converted_at = stamp

    from savdex.adminsite import SavdexModelAdmin

    audit.record(
        connection,
        action="created",
        section="leads",
        actor=staff,
        subject_type="App\\Models\\Crm\\Lead",
        subject_id=lead.pk,
        subject_label=lead.title,
        changes={"after": SavdexModelAdmin.attributes(lead)},
        note=f"Из базы потенциальных клиентов №{prospect.pk}",
        ip=ip,
    )
    audit.record(
        connection,
        action="updated",
        section=SECTION,
        actor=staff,
        subject_type=MODEL,
        subject_id=prospect.pk,
        subject_label=prospect.label(),
        changes={"before": {"lead_id": None}, "after": {"lead_id": lead.pk}},
        note=f"Перенесён в лиды: лид №{lead.pk}",
        ip=ip,
    )

    return lead
